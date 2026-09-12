"""Preserved scholarly artifacts and mechanical extraction; no model calls."""
from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit, quote, unquote
from uuid import uuid4


class ServiceError(Exception):
    def __init__(self, code, message, *, status="error", **details):
        super().__init__(message)
        self.result = {"status": status, "code": code, "message": message, **details}


def xml_tree(data):
    from lxml import etree
    try:
        root = etree.fromstring(data, parser=etree.XMLParser(resolve_entities=False, load_dtd=False, no_network=True, strip_cdata=False))
        tree = root.getroottree()
    except etree.XMLSyntaxError:
        raise ServiceError("invalid_xml", "The XML document is not well formed") from None
    dtd = tree.docinfo.internalDTD
    if dtd is not None and any(dtd.iterentities()) or any(isinstance(node, etree._Entity) for node in root.iter()):
        raise ServiceError("xml_entities_not_supported", "XML entity declarations and references are not expanded")
    return tree


def xml_media(media):
    return media in ('application/xml','text/xml') or media.endswith('+xml')


class Sanitizer:
    sensitive = re.compile(r"(?:api[-_]?key|password|passwd\d*|authorization|cookie|secret|token|ticket|csrf|signature|session(?:id)?|__cf_chl|^sid$|^_[as]$)", re.I)
    secret_fields = {"apikey", "password", "passwd", "passwd2", "authorization", "cookie", "setcookie", "secret", "clientsecret", "secretkey", "token", "accesstoken", "refreshtoken", "authtoken", "ticket", "csrf", "csrftoken", "xcsrftoken", "xsrftoken", "signature", "sid", "session", "sessionid", "jsessionid", "xapikey", "xelsapikey", "xelsinsttoken", "xamzcredential", "xamzsignature", "xamzsecuritytoken"}

    def __init__(self, secrets=()):
        self.secrets = tuple(sorted({str(s) for s in secrets if s}, key=len, reverse=True))

    def text(self, text):
        for secret in self.secrets:
            for variant in {secret, quote(secret, safe=""), quote(quote(secret, safe=""), safe="")}:
                text = text.replace(variant, "[REDACTED]")
        return re.sub(r"https?://[^\s<>\"'{}]+", lambda match: self._url(match.group()), text)

    def url(self, url):
        return self.text(self._url(url))

    def xml(self, data):
        from lxml import etree
        tree = xml_tree(data)
        for node in tree.iter():
            for name, value in list(node.attrib.items()):
                local_name = name.rsplit("}", 1)[-1]
                node.attrib[name] = self.value({local_name: value})[local_name]
            for field in ("text", "tail"):
                value = getattr(node, field)
                if value:
                    cleaned = self.text(value)
                    if cleaned != value:
                        setattr(node, field, cleaned)
        return etree.tostring(tree, encoding="utf-8", xml_declaration=True)

    def _authentication_parameter(self,name,parts):
        normalized=re.sub(r'[^a-z0-9]','',name.lower())
        return normalized in self.secret_fields or bool(self.sensitive.search(name)) or (name.lower()=='sig' and (parts.hostname or '').lower().endswith('.blob.core.windows.net'))

    def _url(self, url):
        original = url
        url = url.replace("&amp;", "&")
        try:
            parts = urlsplit(url)
            port = parts.port
        except ValueError:
            return '[REDACTED_INVALID_URL]'
        path = parts.path
        # Some download actions embed a signed URL as a Tapestry/percent-encoded path argument.
        for match in re.finditer(r'https?(?::|\$3a|%3a)(?:/|\$2f|%2f){2}', path, re.I):
            nested = re.sub(r'\$([0-9a-f]{2})', lambda m: chr(int(m[1], 16)), path[match.start():], flags=re.I)
            try:
                nested = urlsplit(unquote(nested))
                nested.port
            except ValueError:
                path = path[:match.start()] + '[REDACTED_INVALID_URL]'
                break
            if nested.username or nested.password or any(self._authentication_parameter(k,nested) for k, _ in parse_qsl(nested.query, keep_blank_values=True)):
                path = path[:match.start()] + '[REDACTED_DOWNLOAD_URL]'
                break
        # Credentials in userinfo, fragments and authentication query fields are never evidence.
        host = parts.hostname or ""
        if port:
            host += ":" + str(port)
        pairs = parse_qsl(parts.query, keep_blank_values=True)
        query = [(k, self._url(v) if v.startswith(("http://", "https://")) else v)
            for k, v in pairs if not self._authentication_parameter(k,parts)]
        fragment = parts.fragment
        if any(self._authentication_parameter(k,parts) for k, _ in parse_qsl(fragment, keep_blank_values=True)):
            fragment = ""
        if not parts.username and not parts.password and path == parts.path and query == pairs and fragment == parts.fragment:
            return original
        return urlunsplit((parts.scheme, host, path, urlencode(query), fragment))

    def value(self, value):
        if isinstance(value, dict):
            result = {}
            for key, item in value.items():
                normalized = re.sub(r"[^a-z0-9]", "", key.lower())
                public_signature = normalized == "signature" and (item is None or (isinstance(item, str) and item.startswith(("-----BEGIN PGP SIGNATURE-----", "-----BEGIN SSH SIGNATURE-----"))))
                if (normalized in self.secret_fields or normalized.startswith('cfchl') or normalized in {'cupmdtk','csrfmiddlewaretoken'}) and item is not None and not public_signature:
                    result[key] = "[REDACTED]"
                elif isinstance(item, str) and normalized.endswith(("url", "href", "src")):
                    result[key] = self.url(item)
                else:
                    result[key] = self.value(item)
            return result
        if isinstance(value, list):
            return [self.value(v) for v in value]
        if isinstance(value, str):
            if value.startswith(("https://", "http://")):
                return self.url(value)
            return self.text(value)
        return value


def encoded(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def digest(data):
    return hashlib.sha256(data).hexdigest()


class ArtifactStore:
    def __init__(self, root, sanitizer=None):
        self.root = Path(root).resolve()
        self.sanitizer = sanitizer or Sanitizer()

    def _path(self, area, key):
        path = (self.root / area / key).resolve()
        if not path.is_relative_to(self.root):
            raise ServiceError("invalid_reference", "Artifact resolves outside its store")
        return path

    def _write(self, path, data):
        path.parent.mkdir(parents=True, exist_ok=True)
        if path.exists():
            if path.read_bytes() != data:
                raise ServiceError("integrity_error", "Stored artifact bytes do not match their identity")
            return
        temporary = path.with_name("." + uuid4().hex + ".tmp")
        try:
            temporary.write_bytes(data)
            temporary.replace(path)
        finally:
            temporary.unlink(missing_ok=True)

    def put(self, data, media_type, source, url, *, preserve_bytes=False, **context):
        if isinstance(data, str):
            data = data.encode("utf-8")
        if not preserve_bytes and media_type == "application/json":
            data = encoded(self.sanitizer.value(json.loads(data)))
        elif not preserve_bytes and xml_media(media_type) and 'html' not in media_type:
            data = self.sanitizer.xml(data)
            context["xml_normalized"] = True
        elif not preserve_bytes and (media_type.startswith("text/") or media_type in {"application/x-wos", "application/x-bibtex", "application/x-research-info-systems", "application/x-endnote-refer", "application/x-tex"}):
            data = self.sanitizer.text(data.decode("utf-8")).encode("utf-8")
        sha = digest(data)
        manifest = self.sanitizer.value({"sha256": sha, "bytes": len(data), "media_type": media_type,
            "source": source, "url": self.sanitizer.url(url),
            "retrieved_at": datetime.now(timezone.utc).isoformat(), **context})
        raw_manifest = encoded(manifest)
        key = digest(raw_manifest)
        self._write(self._path("blobs", sha), data)
        self._write(self._path("manifests", key), raw_manifest)
        return {"ref": "scholarly://artifact/" + key, **manifest}

    def read(self, ref):
        match = re.fullmatch(r"scholarly://artifact/([0-9a-f]{64})", ref)
        if not match:
            raise ServiceError("invalid_reference", "Use an artifact ref returned by this server")
        key = match.group(1)
        try:
            raw = self._path("manifests", key).read_bytes()
            if digest(raw) != key:
                raise ServiceError("integrity_error", "Artifact manifest has changed")
            manifest = json.loads(raw)
            sha = manifest["sha256"]
            if not re.fullmatch(r"[0-9a-f]{64}", sha):
                raise ServiceError("integrity_error", "Invalid artifact content hash")
            data = self._path("blobs", sha).read_bytes()
            if digest(data) != sha or len(data) != manifest["bytes"]:
                raise ServiceError("integrity_error", "Artifact content has changed")
            return manifest, data
        except FileNotFoundError:
            raise ServiceError("artifact_not_found", "Artifact is not present in this store", status="not_found") from None


def pointer(value, path):
    """RFC 6901 JSON Pointer, with no expression evaluation."""
    if not path:
        return value
    if not path.startswith("/"):
        raise ServiceError("invalid_selection", "JSON selection must be a JSON Pointer starting with /")
    try:
        for part in path[1:].split("/"):
            part = part.replace("~1", "/").replace("~0", "~")
            if isinstance(value, list):
                if not re.fullmatch(r"0|[1-9][0-9]*", part):
                    raise KeyError(part)
                value = value[int(part)]
            else:
                value = value[part]
        return value
    except (KeyError, IndexError, TypeError):
        raise ServiceError("selection_not_found", "Requested JSON location does not exist", status="not_found", selection=path) from None


def bibliography(data, media):
    """Reuse the repository's preserving parsers without invoking their catalog builders."""
    if media == 'application/x-endnote-refer':
        rows, fields, raw_lines, last = [], None, [], None
        for line in data.decode('utf-8-sig').splitlines():
            if line.startswith('%0 '):
                if fields is not None:
                    rows.append({'fields': fields, 'raw_text': '\n'.join(raw_lines).rstrip()})
                fields, raw_lines, last = {}, [], None
            if fields is None:
                if line.strip():
                    raise ServiceError('bibliography_parse_error', 'EndNote record must begin with %0')
                continue
            raw_lines.append(line)
            match = re.match(r'^%(DOI|\S)(?: (.*))?$', line)
            if match:
                last = match[1]
                fields.setdefault(last, []).append(match[2] or '')
            elif line.startswith('%'):
                raise ServiceError('bibliography_parse_error', 'Malformed EndNote field')
            elif line.strip() and last:
                fields[last][-1] += '\n' + line
        if fields is not None:
            rows.append({'fields': fields, 'raw_text': '\n'.join(raw_lines).rstrip()})
        return rows
    if media == "application/x-wos":
        rows, fields, raw_lines, last = [], None, [], None
        for line in data.decode("utf-8-sig").splitlines():
            tag = line[:2]
            if tag == "PT":
                if fields is not None:
                    raise ServiceError("bibliography_parse_error", "WOS record is missing ER")
                fields, raw_lines = {}, []
            if fields is None:
                if line.strip() and tag not in ("FN", "VR", "EF"):
                    raise ServiceError("bibliography_parse_error", "Unexpected text outside a WOS record")
                continue
            raw_lines.append(line)
            if tag == "ER":
                rows.append({"id": fields.get("ut", [None])[0], "fields": fields, "raw_text": "\n".join(raw_lines)})
                fields, last = None, None
            elif re.match(r"^[A-Z0-9]{2} ", line):
                last = tag.lower()
                fields.setdefault(last, []).append(line[3:])
            elif line.startswith("   ") and last:
                fields[last].append(line[3:])
            elif line.strip():
                raise ServiceError("bibliography_parse_error", "Malformed WOS field or continuation")
        if fields is not None:
            raise ServiceError("bibliography_parse_error", "Final WOS record is missing ER")
        return rows
    import importlib.util
    import sys
    script_name = "build_b002_catalog" if media == "application/x-bibtex" else "build_b001_catalog"
    module_name = "scholarly_" + script_name
    if module_name not in sys.modules:
        path = Path(__file__).resolve().parents[2] / "research/pvsnp/literature-baseline/scripts" / (script_name + ".py")
        spec = importlib.util.spec_from_file_location(module_name, path)
        module = importlib.util.module_from_spec(spec)
        sys.modules[module_name] = module
        spec.loader.exec_module(module)
    module = sys.modules[module_name]
    text = data.decode("utf-8-sig")
    try:
        if media == "application/x-bibtex":
            rows = []
            # Native exports can have spaced citation keys; extract them with a warning, without rewriting.
            scanner = module._BibScanner(text, source_name="artifact", allow_key_whitespace=True)
            entries = []
            while scanner.position < len(text):
                scanner.skip_space_and_comments()
                if scanner.position == len(text):
                    break
                if text[scanner.position] != "@":
                    # BibTeX ignores free annotations outside entries; the artifact retains them.
                    scanner.position += 1
                    continue
                entries.append(scanner.read_entry())
            for entry in entries:
                row = {"id": entry.key, "type": entry.entry_type, "raw_text": entry.raw_text, "parse_warnings": list(entry.parse_warnings)}
                for field, value in entry.field_items:
                    if field not in row:
                        row[field] = value
                    elif isinstance(row[field], list):
                        row[field].append(value)
                    else:
                        row[field] = [row[field], value]
                rows.append(row)
            return rows
        first = re.search(r'(?m)^TY  - ',text)
        preamble = text[:first.start()] if first else ''
        if preamble.strip() and all(not line.strip() or line.startswith(('Provider:', 'Database:', 'Content:')) for line in preamble.splitlines()):
            rows = module.parse_ris(text[first.start():], source_name="artifact")
            if rows:
                rows[0]['source_preamble'] = preamble
            return rows
        return module.parse_ris(text, source_name="artifact")
    except ValueError:
        raise ServiceError("bibliography_parse_error", "Bibliographic input is malformed or uses syntax not supported by the preserving parser; original bytes remain in the artifact") from None


def legacy_spreadsheet_rows(data,ref,sheet_name,start,count,fields):
    import xlrd
    from openpyxl.utils import get_column_letter,column_index_from_string
    try:
        workbook=xlrd.open_workbook(file_contents=data,on_demand=True,formatting_info=True)
    except Exception:
        raise ServiceError('invalid_document','The XLS workbook cannot be read') from None
    try:
        sheets=[{'name':sheet.name,'declared_rows':sheet.nrows,'declared_columns':sheet.ncols} for sheet in workbook.sheets()]
        if not sheet_name and len(sheets)>1:
            return {'status':'selection_required','artifact':ref,'sheets':sheets,'message':'Set selection to an exact worksheet name'}
        if sheet_name and sheet_name not in workbook.sheet_names():
            raise ServiceError('selection_not_found','Worksheet name is not present',status='not_found',sheets=sheets)
        if not sheets:
            raise ServiceError('no_tabular_worksheets','This workbook contains no tabular worksheet',status='unsupported')
        if fields and any(not isinstance(name,str) or not re.fullmatch(r'[A-Z]{1,3}',name) or column_index_from_string(name)>16384 for name in fields):
            raise ServiceError('invalid_selection','Spreadsheet fields are column letters such as A, B or AA')
        sheet=workbook.sheet_by_name(sheet_name) if sheet_name else workbook.sheet_by_index(0)
        items=[]
        for index in range(start,min(start+count,sheet.nrows)):
            values,formats,types={},{},{}
            for col in range(sheet.ncols):
                letter=get_column_letter(col+1)
                if fields and letter not in fields:
                    continue
                cell=sheet.cell(index,col)
                value=cell.value
                if cell.ctype==xlrd.XL_CELL_DATE:
                    value=xlrd.xldate_as_datetime(value,workbook.datemode).isoformat()
                    types[letter]='d'
                elif cell.ctype==xlrd.XL_CELL_BOOLEAN:
                    value=bool(value)
                elif cell.ctype in (xlrd.XL_CELL_EMPTY,xlrd.XL_CELL_BLANK):
                    value=None
                elif cell.ctype==xlrd.XL_CELL_ERROR:
                    value=xlrd.error_text_from_code.get(value,str(value))
                    types[letter]='e'
                values[letter]=value
                xf=workbook.xf_list[sheet.cell_xf_index(index,col)]
                format=workbook.format_map[xf.format_key].format_str
                if format!='General':
                    formats[letter]=format
            item={'row':index+1,'values':values}
            if fields and any(name not in values for name in fields):
                item['missing_columns']=[name for name in fields if name not in values]
            if formats:
                item['number_formats']=formats
            if types:
                item['cell_types']=types
            items.append(item)
        return {'status':'ok','artifact':ref,'unit':'spreadsheet_rows','sheet':sheet.name,'sheets':sheets,'start':start,
                'total':sheet.nrows,'returned':len(items),'next_start':start+len(items) if start+len(items)<sheet.nrows else None,
                'items':items,'interpretation':'Stored XLS values and number formats; formula definitions, embedded objects and hyperlinks remain in the original file; nothing is executed'}
    finally:
        workbook.release_resources()


def spreadsheet_rows(data,ref,sheet_name,start,count,fields):
    import io,zipfile
    from datetime import date,time
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter,column_index_from_string
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist())>256*1024*1024:
                raise ServiceError('resource_limit','Uncompressed workbook exceeds the 256 MiB extraction limit')
        workbook=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
    except ServiceError:
        raise
    except Exception:
        raise ServiceError('invalid_document','The XLSX workbook cannot be read') from None
    try:
        sheets=[{'name':sheet.title,'declared_rows':sheet.max_row,'declared_columns':sheet.max_column,'state':sheet.sheet_state} for sheet in workbook.worksheets]
        if not sheets:
            raise ServiceError('no_tabular_worksheets','This workbook contains no tabular worksheet',status='unsupported')
        if not sheet_name and len(sheets)>1:
            return {'status':'selection_required','artifact':ref,'sheets':sheets,'message':'Set selection to an exact worksheet name'}
        if sheet_name and sheet_name not in workbook.sheetnames:
            raise ServiceError('selection_not_found','Worksheet name is not present',status='not_found',sheets=sheets)
        sheet=workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        if fields:
            try:
                if any(not isinstance(name,str) or not re.fullmatch(r'[A-Z]{1,3}',name) or column_index_from_string(name)>16384 for name in fields):
                    raise ValueError
            except ValueError:
                raise ServiceError('invalid_selection','Spreadsheet fields are column letters such as A, B or AA') from None
        # Source worksheet dimensions can be stale; stream the actual rows to establish the range.
        sheet.reset_dimensions()
        items=[]
        total=0
        for number,row in enumerate(sheet.iter_rows(),1):
            total=number
            if not start<number<=start+count:
                continue
            values,formats,formulas,types={},{},[],{}
            for index,cell in enumerate(row,1):
                column=get_column_letter(index)
                if fields and column not in fields:
                    continue
                value=cell.value
                if isinstance(value,(datetime,date,time)):
                    value=value.isoformat()
                elif value is not None and not isinstance(value,(str,int,float,bool)):
                    raise ServiceError('cell_type_not_supported','This workbook cell type requires a separate reader',cell=f'{sheet.title}!{column}{number}')
                values[column]=value
                if cell.number_format!='General':
                    formats[column]=cell.number_format
                if cell.data_type=='f':
                    formulas.append(column)
                if cell.data_type in ('f','e','d'):
                    types[column]=cell.data_type
            item={'row':number,'values':values}
            if fields and any(name not in values for name in fields):
                item['missing_columns']=[name for name in fields if name not in values]
            if formats:
                item['number_formats']=formats
            if formulas:
                item['formula_columns']=formulas
            if types:
                item['cell_types']=types
            items.append(item)
        return {'status':'ok','artifact':ref,'unit':'spreadsheet_rows','sheet':sheet.title,'sheets':sheets,'start':start,
                'total':total,'returned':len(items),'next_start':start+len(items) if start+len(items)<total else None,
                'items':items,'interpretation':'Stored cell values and number formats; formulas are returned as text and are not executed'}
    finally:
        workbook.close()


def extract_artifact(store, ref, *, fields=None, start=0, count=20, selection="", render=False):
    if type(start) is not int or start < 0 or type(count) is not int or count < 1:
        raise ServiceError("invalid_range", "start must be nonnegative and count must be positive")
    meta, data = store.read(ref)
    media = meta["media_type"]
    unit, missing = "records", []
    if media in {"image/png", "image/jpeg", "image/webp", "image/gif"}:
        if selection or fields or start:
            raise ServiceError("invalid_selection", "An image ref is read as one whole image")
        return {"status": "ok", "artifact": ref, "unit": "image", "image_ref": ref, "media_type": media, "bytes": len(data)}
    if media == "application/json":
        value = pointer(json.loads(data), selection)
        rows = value if isinstance(value, list) else [value]
    elif media in {"application/zip", "application/x-zip-compressed", "application/gzip", "application/x-gzip", "application/x-tar"}:
        import io, mimetypes, tarfile, zipfile
        is_zip = zipfile.is_zipfile(io.BytesIO(data))
        archive = zipfile.ZipFile(io.BytesIO(data)) if is_zip else tarfile.open(fileobj=io.BytesIO(data), mode="r:*")
        with archive:
            members = archive.infolist() if is_zip else archive.getmembers()
            rows = [{"path": m.filename if is_zip else m.name, "bytes": m.file_size if is_zip else m.size,
                "directory": m.is_dir() if is_zip else m.isdir()} for m in members]
            if selection:
                matches = [(m, row) for m, row in zip(members, rows) if row["path"] == selection]
                if len(matches) != 1:
                    raise ServiceError("selection_not_found", "Archive member must exist exactly once", status="not_found", selection=selection)
                member, row = matches[0]
                if row["directory"] or (not is_zip and not member.isfile()):
                    raise ServiceError("invalid_selection", "Select a regular archive file")
                if row["bytes"] > 256 * 1024 * 1024:
                    raise ServiceError("resource_limit", "Selected archive member exceeds the 256 MiB extraction limit; the archive remains preserved")
                content = archive.read(member) if is_zip else archive.extractfile(member).read()
                suffix = Path(selection).suffix.lower()
                mime = {".bib": "application/x-bibtex", ".ris": "application/x-research-info-systems"}.get(suffix) or mimetypes.guess_type(selection)[0]
                if mime is None:
                    try:
                        content.decode("utf-8")
                        mime = "text/plain" if b"\x00" not in content else "application/octet-stream"
                    except UnicodeDecodeError:
                        mime = "application/octet-stream"
                child = store.put(content, mime, meta["source"], meta["url"], preserve_bytes=True, parent_ref=ref, archive_member=selection)
                rows = [{**row, "artifact": child}]
        unit = "archive_members"
    elif media in {"text/csv", "application/x-research-info-systems", "application/x-bibtex", "application/x-wos", "application/x-endnote-refer"}:
        if selection == "raw":
            rows = data.decode("utf-8-sig").splitlines()
            unit = "lines"
        elif media == "text/csv":
            import csv, io
            rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
        else:
            rows = bibliography(data, media)
    elif "html" in media:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(data, "html.parser")
        if selection:
            try:
                rows = [str(node) for node in soup.select(selection)]
            except Exception:
                raise ServiceError("invalid_selection", "Invalid HTML CSS selector") from None
        else:
            rows = data.decode("utf-8").splitlines()
            unit = "lines"
    elif media == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
        return spreadsheet_rows(data,ref,selection,start,count,fields)
    elif media == 'application/vnd.ms-excel':
        return legacy_spreadsheet_rows(data,ref,selection,start,count,fields)
    elif xml_media(media):
        from lxml import etree
        try:
            tree = xml_tree(data)
            rows = [etree.tostring(node, encoding="unicode") for node in tree.getroot().findall(selection)] if selection else [etree.tostring(tree, encoding="unicode")]
        except ServiceError:
            raise
        except Exception:
            raise ServiceError("invalid_selection", "XML could not be parsed or the XPath subset is invalid") from None
        unit = "elements"
    elif media == "application/pdf":
        import pymupdf
        try:
            doc = pymupdf.open(stream=data, filetype="pdf")
        except Exception:
            raise ServiceError("invalid_document", "PDF cannot be opened") from None
        with doc:
            total = len(doc)
            rows = []
            for index in range(start, min(start + count, total)):
                page = doc[index]
                text = page.get_text(sort=True)
                row = {"page": index + 1, "text": text, "text_layer_empty": not text.strip()}
                if render or not text.strip():
                    png = page.get_pixmap(dpi=144).tobytes("png")
                    row["image"] = store.put(png, "image/png", meta["source"], meta["url"], parent_ref=ref, page=index + 1)
                rows.append(row)
        return {"status": "ok", "artifact": ref, "unit": "pages", "start": start, "total": total,
            "returned": len(rows), "next_start": start + len(rows) if start + len(rows) < total else None,
            "items": rows, "interpretation": "Mechanical text and rendering; image meaning and text accuracy are assessed by the caller"}
    elif media.startswith("text/") or media in ('application/x-tex', 'application/x-latex'):
        rows = data.decode("utf-8").splitlines()
        unit = "lines"
    else:
        raise ServiceError("format_not_supported", "Use the preserved binary artifact; this format has no mechanical extractor", status="unsupported", artifact=ref)
    total = len(rows)
    window = rows[start:start + count]
    if fields:
        selected = []
        for index, row in enumerate(window, start):
            if not isinstance(row, dict):
                raise ServiceError("invalid_selection", "fields requires object records")
            absent = [field for field in fields if field not in row]
            if absent:
                missing.append({"index": index, "fields": absent})
            selected.append({field: row[field] for field in fields if field in row})
        window = selected
    return {"status": "ok", "artifact": ref, "unit": unit, "start": start, "total": total,
        "returned": len(window), "next_start": start + len(window) if start + len(window) < total else None,
        "items": window, "missing_fields": missing}
