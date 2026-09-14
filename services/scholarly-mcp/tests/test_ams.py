import json,sys,tempfile,unittest
import io
from pathlib import Path
from unittest.mock import Mock,patch
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))

class AMSContract(unittest.TestCase):
    def test_provider_error_string_is_a_typed_response_failure(self):
        from ams import product_records
        from scholarly import ServiceError
        with self.assertRaises(ServiceError) as raised:
            product_records('{"message":"SYNTHETIC_PROVIDER_ERROR"}')
        self.assertEqual(raised.exception.result['code'],'invalid_response')
        self.assertNotIn('SYNTHETIC_PROVIDER_ERROR',str(raised.exception.result))

    def test_native_zero_count_allows_null_products_without_changing_evidence(self):
        from ams import product_records
        from scholarly import ServiceError
        payload={'Products':None,'Filters':None,'TotalRows':0,'PageNumber':1,'RowsPerPage':12,'SortBy':None}
        self.assertEqual(product_records(payload),([],0,0))
        self.assertIsNone(payload['Products'])
        payload['TotalRows']=1
        with self.assertRaises(ServiceError):product_records(payload)

    def test_book_review_display_title_is_normalized_without_mutating_native_payload(self):
        from ams import product_records
        review={'Title':'','DocumentType':'BOOKREV','BookReviews':[{'Title':'Complexity theory of real functions'}]}
        payload={'Products':[review],'TotalRows':1,'PageNumber':1,'RowsPerPage':12}
        rows,total,first=product_records(payload)
        self.assertEqual((total,first),(1,0))
        self.assertEqual(rows[0]['Title'],'Complexity theory of real functions')
        self.assertEqual(rows[0]['TitleSource'],'BookReviews[0].Title')
        self.assertEqual(payload['Products'][0]['Title'],'')
        self.assertNotIn('TitleSource',payload['Products'][0])

    def test_translated_challenge_is_not_treated_as_a_ready_search_page(self):
        from ams import validate_search_page
        from scholarly import ServiceError
        raw=(ROOT/'tests/fixtures/ams-challenge-ko.html').read_text('utf-8')
        with self.assertRaises(ServiceError) as error:validate_search_page(raw,200)
        self.assertEqual(error.exception.result['status'],'restricted')
        self.assertEqual(error.exception.result['code'],'publisher_access_challenge')

    def test_nested_expression_preserves_native_field_and_operator_meanings(self):
        from ams import advanced_expression,check_expression_request
        from urllib.parse import urlencode,quote
        expression={'condition':'AND','terms':[{'field':'Title','operator':'contains','value':'algebra'},
            {'condition':'OR','terms':[{'field':'PublicationYear','operator':'equals','value':'2000'},
                                     {'field':'PublicationYear','operator':'equals','value':'2001'}]}]}
        native=advanced_expression(expression)
        self.assertEqual(native,json.loads((ROOT/'tests/fixtures/ams-dynamic-query.json').read_text('utf-8')))
        for key in ('advancedSearchJson','dynamicAdvancedSearch'):
            for value in (json.dumps(native),quote(json.dumps(native))):
                check_expression_request(urlencode({key:value}),native)

    def test_invalid_and_changed_expression_are_rejected_before_search(self):
        from ams import advanced_expression,check_expression_request,AMS
        from scholarly import ArtifactStore,ServiceError
        from urllib.parse import urlencode
        leaf={'field':'PublicationYear','operator':'equals','value':'2000'}
        invalid=[{'condition':'NOT','terms':[leaf]}, {'condition':'AND','terms':[]},
                 {'field':'PublicationYear','operator':'contains','value':'2000'},
                 {'field':'Title','operator':'equals','value':''}, {**leaf,'unknown':True}]
        for value in invalid:
            with self.assertRaises(ServiceError):advanced_expression(value)
        native=advanced_expression({'condition':'AND','terms':[leaf]})
        changed={'condition':'AND','queryElements':[{**native['queryElements'][0],'value':'2001'}]}
        for body in (urlencode({'advancedSearchJson':json.dumps(changed)}), 'advancedSearch=ProductTitle%3Dalgebra'):
            with self.assertRaises(ServiceError) as error:check_expression_request(body,native)
            self.assertEqual(error.exception.result['code'],'query_mismatch')
        with tempfile.TemporaryDirectory() as temp:
            adapter=AMS(ArtifactStore(temp),None)
            for query,options in [('algebra',{'expression':{'condition':'AND','terms':[leaf]}}),
                                  ('',{'expression':{'condition':'AND','terms':[leaf]},'author':'Knuth'})]:
                with self.assertRaises(ServiceError) as error:adapter.search(query,options=options)
                self.assertEqual(error.exception.result['code'],'invalid_options')

    def test_or_expression_is_not_replaced_by_a_native_full_title_lookup(self):
        from ams import basic_expression
        from scholarly import ServiceError
        # Actual UI: two Title words gave 915 results; adding OR words gave one
        # exact-title result. Neither spaces nor explicit OR preserved a union.
        for words in ('Quantum realization', 'Quantum realization of the finite element method'):
            expression={'condition':'OR','terms':[
                {'field':'Title','operator':'contains','value':word} for word in words.split()]}
            with self.subTest(words=words):
                with self.assertRaises(ServiceError) as error:
                    basic_expression(expression)
                self.assertEqual(error.exception.result['status'],'unsupported')
                self.assertEqual(error.exception.result['stage'],'search_options')

    def test_verified_basic_expression_subset_preserves_boolean_meaning(self):
        from ams import basic_expression,basic_search_expected,basic_search_request,basic_search_value
        from urllib.parse import urlencode
        title={'field':'Title','operator':'contains','value':'algebra'}
        author={'field':'Author','operator':'contains','value':'smith'}
        self.assertEqual(basic_expression(title),{'Title':'algebra'})
        self.assertEqual(basic_expression({'condition':'AND','terms':[title,author]}),
                         {'Title':'algebra','Author':'smith'})
        self.assertEqual(basic_expression({'field':'Title','operator':'contains','value':'Go\u0308del'}),
                         {'Title':'Gödel'})
        self.assertEqual(basic_search_value('Gödel'), 'Godel')
        self.assertEqual(basic_search_value('Schrödinger'), 'Schrodinger')
        self.assertEqual(basic_search_value('Gödel’s theorem'), "Godel's theorem")
        self.assertEqual(basic_search_expected({'Title':'Gödel'}, {'Title':('control','ProductTitle')}),
                         {'ProductTitle':['Godel']})
        body=urlencode({'advancedSearch':'ProductTitle=algebra|Author=smith','pageNumber':'1'})
        self.assertEqual(basic_search_request(body),{'ProductTitle':['algebra'],'Author':['smith']})
        implicit=urlencode({'advancedSearch':'ProductTitle=algebra topology geometry','pageNumber':'1'})
        self.assertEqual(basic_search_request(implicit),{'ProductTitle':['algebra topology geometry']})

    def test_singleton_and_distinct_field_groups_preserve_boolean_meaning(self):
        from ams import basic_expression
        title=lambda value:{'field':'Title','operator':'contains','value':value}
        author=lambda value:{'field':'Author','operator':'contains','value':value}
        nested={'condition':'OR','terms':[{'condition':'AND','terms':[title('algebra')]}]}
        self.assertEqual(basic_expression(nested),{'Title':'algebra'})
        combined={'condition':'AND','terms':[nested,author('smith')]}
        self.assertEqual(basic_expression(combined),{'Title':'algebra','Author':'smith'})
        self.assertEqual(basic_expression({'condition':'OR','terms':[combined]}),
                         {'Title':'algebra','Author':'smith'})
        self.assertEqual(basic_expression({'condition':'OR','terms':[author('Boykan'),author('Jech')]}),
                         {'Author':'Boykan Jech'})

    def test_unverified_expression_shapes_are_rejected_before_browser_navigation(self):
        from ams import AMS,basic_expression
        from scholarly import ArtifactStore,ServiceError
        title={'field':'Title','operator':'contains','value':'algebra'}
        provider_unsupported=[
            {'condition':'AND','terms':[title,{**title,'value':'topology'}]},
            {'condition':'OR','terms':[title,{'field':'Author','operator':'contains','value':'smith'}]},
            {'condition':'OR','terms':[title,{**title,'value':'topology'}]},
        ]
        locally_unmapped=[
            {'field':'PublicationYear','operator':'equals','value':'2000'},
            {'field':'Title','operator':'equals','value':'algebra'},
            {'field':'Title','operator':'contains','value':'algebra topology'},
            {'field':'Title','operator':'contains','value':'OR'},
        ]
        for expression in provider_unsupported:
            with self.assertRaises(ServiceError) as error:
                basic_expression(expression)
            self.assertEqual(error.exception.result['status'],'unsupported')
            self.assertEqual(error.exception.result['code'],'provider_feature_unavailable')
        for expression in locally_unmapped:
            with self.assertRaises(ServiceError) as error:
                basic_expression(expression)
            self.assertEqual(error.exception.result['status'],'unsupported')
            self.assertEqual(error.exception.result['code'],'expression_mapping_unavailable')
        with tempfile.TemporaryDirectory() as temp:
            gateway=Mock()
            adapter=AMS(ArtifactStore(temp),gateway)
            with self.assertRaises(ServiceError) as error:
                adapter.search('',options={'expression':provider_unsupported[0]})
        self.assertEqual(error.exception.result['status'],'unsupported')
        self.assertEqual(error.exception.result['code'],'provider_feature_unavailable')
        gateway.open.assert_not_called()

    def test_verified_expression_uses_the_direct_official_search_page(self):
        from ams import AMS
        from scholarly import ArtifactStore
        expression={'field':'Title','operator':'contains','value':'algebra'}
        with tempfile.TemporaryDirectory() as temp:
            gateway=Mock()
            gateway.public_page.side_effect=RuntimeError('route selected')
            adapter=AMS(ArtifactStore(temp),gateway)
            with self.assertRaisesRegex(RuntimeError,'route selected'):
                adapter.search('',options={'expression':expression})
        gateway.public_page.assert_called_once_with('ams-expression-public')
        gateway.open.assert_not_called()

    def test_each_expression_search_discards_only_its_previous_public_page(self):
        from ams import AMS
        from scholarly import ArtifactStore
        expression={'field':'Title','operator':'contains','value':'algebra'}
        cached=Mock();cached.is_closed.return_value=False
        institutional=Mock()
        with tempfile.TemporaryDirectory() as temp:
            gateway=Mock();gateway.pages={'ams-expression-public':cached,'ams':institutional}
            gateway.public_page.side_effect=RuntimeError('fresh public page requested')
            adapter=AMS(ArtifactStore(temp),gateway)
            with self.assertRaisesRegex(RuntimeError,'fresh public page requested'):
                adapter.search('',options={'expression':expression})
        cached.close.assert_called_once_with()
        self.assertIs(gateway.pages['ams'],institutional)
        self.assertNotIn('ams-expression-public',gateway.pages)

    def test_expression_cursor_is_rejected_before_public_page_creation(self):
        from ams import AMS
        from scholarly import ArtifactStore,ServiceError
        expression={'field':'Title','operator':'contains','value':'algebra'}
        with tempfile.TemporaryDirectory() as temp:
            gateway=Mock()
            adapter=AMS(ArtifactStore(temp),gateway)
            with self.assertRaises(ServiceError) as error:
                adapter.search('',options={'expression':expression},cursor='opaque')
        self.assertEqual(error.exception.result['code'],'invalid_options')
        gateway.public_page.assert_not_called()
        gateway.open.assert_not_called()

    def test_expression_validation_keeps_range_export_and_option_boundaries(self):
        from ams import AMS
        from scholarly import ArtifactStore,ServiceError
        expression={'field':'Title','operator':'contains','value':'algebra'}
        with tempfile.TemporaryDirectory() as temp:
            gateway=Mock();adapter=AMS(ArtifactStore(temp),gateway)
            cases=[({'start':-1},'invalid_range'),({'limit':0},'invalid_range'),
                   ({'export':'ris'},'invalid_export')]
            for arguments,code in cases:
                with self.subTest(code=code),self.assertRaises(ServiceError) as error:
                    adapter.search('',options={'expression':expression},**arguments)
                self.assertEqual(error.exception.result['code'],code)
            with self.assertRaises(ServiceError) as error:
                adapter.search('',options={'expression':expression,'author':'smith'})
            self.assertEqual(error.exception.result['code'],'invalid_options')
        gateway.public_page.assert_not_called();gateway.open.assert_not_called()

    def test_basic_search_preserves_the_gateway_arrival_host(self):
        from ams import AMS
        from scholarly import ArtifactStore
        page=Mock();page.url='https://pubs-ams-org.libproxy.knu.ac.kr/landing'
        with tempfile.TemporaryDirectory() as temp:
            gateway=Mock();gateway.open.return_value=page
            adapter=AMS(ArtifactStore(temp),gateway)
            with patch('ams.goto_source',side_effect=RuntimeError('route observed')) as go:
                with self.assertRaisesRegex(RuntimeError,'route observed'):
                    adapter.search('algebra',options={})
        gateway.open.assert_called_once_with('ams','AMS',adapter.profile.hosts,'body')
        gateway.public_page.assert_not_called()
        self.assertEqual(go.call_args.args[1],
                         'https://pubs-ams-org.libproxy.knu.ac.kr/AdvancedSearch')

    def test_source_guide_offers_only_the_verified_expression_subset(self):
        from registry import SOURCES
        guide=SOURCES['ams'].usage
        self.assertIn('expression',guide)
        self.assertIn('Title contains',guide['expression'])
        self.assertIn('Unicode',guide['expression'])
        self.assertIn('expression_mapping_unavailable',guide['options'])
        examples=[example.get('arguments',{}).get('options',{}).get('expression')
                  for example in guide['examples']]
        from ams import basic_expression
        expressions=[value for value in examples if value]
        self.assertTrue(expressions)
        for expression in expressions:
            self.assertTrue(basic_expression(expression))

    def test_unavailable_search_page_is_not_an_input_timeout_or_no_results(self):
        from ams import validate_search_page
        from scholarly import ServiceError
        raw=(ROOT/'tests/fixtures/ams-unavailable.html').read_text('utf-8')
        with self.assertRaises(ServiceError) as cm:validate_search_page(raw)
        self.assertEqual(cm.exception.result['code'],'provider_unavailable')
        self.assertEqual(cm.exception.result['status'],'error')
        self.assertTrue(cm.exception.result['retryable'])
        with self.assertRaises(ServiceError) as cm:validate_search_page('<html><body>Denied</body></html>',403)
        self.assertEqual(cm.exception.result['status'],'restricted')
        validate_search_page('<html><head><title>Advanced Search</title></head><body>Search</body></html>',200)

    def test_product_code_url_does_not_require_a_slash_in_the_code(self):
        from ams import AMS
        from scholarly import ArtifactStore
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as temp:
            adapter=AMS(ArtifactStore(temp),None)
            target='https://pubs.ams.org/view?ProductCode=MPRIZE'
            row={'ProductCode':'MPRIZE','Title':'The Millennium Prize Problems','target':target}
            proof=adapter.store.put(json.dumps([row]),'application/json','ams',target)
            def lookup(query,**kwargs):
                adapter.known['mprize']=(row,proof)
                return {'total':1,'received':1}
            with patch.object(adapter,'search',side_effect=lookup) as search,patch.object(adapter,'document',side_effect=AssertionError('Expected catalogue lookup')):
                found=adapter.fetch(target,parts=['metadata'])
            self.assertEqual(found['status'],'ok')
            self.assertEqual(search.call_args.args[0],'MPRIZE')

    def test_mixed_products_keep_subtitle_article_doi_and_meeting_abstract(self):
        from ams import product_records
        payload=json.loads((ROOT/'tests/fixtures/ams-products.json').read_text('utf-8'))
        rows,total,first=product_records(payload)
        self.assertEqual((len(rows),total,first),(12,6587,0))
        self.assertEqual(rows[0]['ProductCode'],'SURV/168')
        self.assertIn('Lie Algebras',rows[0]['Subtitle'])
        self.assertEqual(rows[2]['DOI'],'10.1090/jams/1045')
        self.assertEqual(rows[3]['ContentType'],'ABS')
        self.assertIsNone(rows[3]['DOI'])

    def test_over_capacity_and_empty_intermediate_are_range_errors(self):
        from ams import product_records
        from scholarly import ServiceError
        payload=json.loads((ROOT/'tests/fixtures/ams-products.json').read_text('utf-8'))
        payload['Products'].append(dict(payload['Products'][0]))
        with self.assertRaises(ServiceError) as over:
            product_records(payload)
        self.assertEqual(over.exception.result['code'],'range_mismatch')
        payload={'Products':[],'TotalRows':24,'PageNumber':2,'RowsPerPage':12}
        with self.assertRaises(ServiceError) as empty:
            product_records(payload)
        self.assertEqual(empty.exception.result['code'],'range_mismatch')

    def test_provider_short_positive_page_is_a_valid_native_prefix(self):
        from ams import product_records
        payload={'Products':[{'Title':f'Product {index}','ProductId':str(index)} for index in range(11)],
                 'TotalRows':8685,'PageNumber':1,'RowsPerPage':12}
        rows,total,first=product_records(payload)
        self.assertEqual((len(rows),total,first),(11,8685,0))

    def test_page_sequence_rejects_changed_page_number_or_total(self):
        from ams import validate_page_sequence
        from scholarly import ServiceError
        validate_page_sequence(2,2,8685,8685)
        for observed_page,reported in ((3,8685),(2,8684)):
            with self.assertRaises(ServiceError) as error:
                validate_page_sequence(observed_page,2,8685,reported)
            self.assertEqual(error.exception.result['code'],'range_mismatch')

    def test_short_last_page_is_valid_but_an_intersecting_ordinal_gap_is_partial(self):
        from ams import product_records,requested_page_gap
        last={'Products':[{'Title':f'Last {index}'} for index in range(9)],
              'TotalRows':8685,'PageNumber':724,'RowsPerPage':12}
        rows,total,first=product_records(last)
        self.assertEqual((len(rows),total,first),(9,8685,8676))
        self.assertIsNone(requested_page_gap(first,len(rows),12,total,8676,9))
        gap=requested_page_gap(0,11,12,8685,0,20)
        self.assertEqual((gap['code'],gap['missing_start'],gap['missing_end'],gap['resume_start']),
                         ('source_range_gap',11,12,12))
        self.assertIsNone(requested_page_gap(0,11,12,8685,0,11))
        self.assertIsNone(requested_page_gap(0,11,12,8685,12,20))

    def test_page_export_scope_does_not_claim_a_complete_gapped_window(self):
        from ams import page_export_scope
        for format in ('native','html'):
            complete=page_export_scope(format,False)
            partial=page_export_scope(format,True)
            self.assertIn('requested window',complete)
            self.assertNotIn('complete requested window',partial)
            self.assertIn('confirmed prefix',partial)
            self.assertIn('window is incomplete',partial)

    def test_product_id_collision_is_not_a_duplicate_but_identical_records_are(self):
        from ams import native_record_fingerprint,select_native_page
        first={'ProductId':73989,'ProductCode':'ABS','Title':'Topology-Enhanced Deep Learning'}
        second={'ProductId':73989,'ProductCode':'ABS','Title':'Bridging Climate Science'}
        self.assertNotEqual(native_record_fingerprint(first),native_record_fingerprint(second))
        self.assertEqual(native_record_fingerprint(first),native_record_fingerprint(dict(first)))
        chosen,issue=select_native_page([first,second],0,2,2,0,2,set())
        self.assertEqual(chosen,[(0,first),(1,second)])
        self.assertIsNone(issue)

    def test_identical_native_duplicate_preserves_prefix_and_resume_ordinal(self):
        from ams import select_native_page
        rows=[{'Title':'First','ProductId':1},{'Title':'Repeated','ProductId':2},
              {'Title':'Repeated','ProductId':2},{'Title':'After','ProductId':3}]
        chosen,issue=select_native_page(rows,0,4,4,0,4,set())
        self.assertEqual(chosen,[(0,rows[0]),(1,rows[1])])
        self.assertEqual((issue['code'],issue['duplicate_at'],issue['resume_start']),
                         ('source_duplicate_record',2,3))

    def test_native_xlsx_keeps_variant_rows_and_binds_one_logical_product(self):
        from ams import validate_product_xlsx
        from openpyxl import Workbook
        row=json.loads((ROOT/'tests/fixtures/ams-products.json').read_text('utf-8'))['Products'][0]
        workbook=Workbook();sheet=workbook.active;sheet.title='Products'
        sheet.append(['Product Code','Title','ISBN','Author/Editor','Publication Date'])
        for isbn in ('978-0-8218-5262-0','978-1-4704-1395-8',''):
            sheet.append(['','Algebras, Rings and Modules: Lie Algebras and Hopf Algebras',isbn,
                          'Michiel Hazewinkel; Nadiya Gubareni; V. V. Kirichenko',2010])
        output=io.BytesIO();workbook.save(output);workbook.close()
        found=validate_product_xlsx(output.getvalue(),row)
        self.assertEqual(found['physical_rows'],3)
        self.assertEqual(found['logical_products'],1)
        self.assertEqual(found['product_code'],'SURV/168')
        self.assertEqual(found['product_code_cells'],0)
        self.assertEqual(found['native_isbn_rows'],1)

    def test_native_xlsx_rejects_another_product_without_filling_blank_codes(self):
        from ams import validate_product_xlsx
        from openpyxl import Workbook
        from scholarly import ServiceError
        row=json.loads((ROOT/'tests/fixtures/ams-products.json').read_text('utf-8'))['Products'][0]
        workbook=Workbook();sheet=workbook.active;sheet.append(['Product Code','Title','ISBN','Author/Editor','Publication Date'])
        sheet.append(['','Another Product','978-0-8218-5262-0','Different Author',2010])
        output=io.BytesIO();workbook.save(output);workbook.close()
        with self.assertRaises(ServiceError) as error:validate_product_xlsx(output.getvalue(),row)
        self.assertEqual(error.exception.result['code'],'identity_mismatch')

    def test_native_xlsx_rejects_nonblank_code_from_another_product(self):
        from ams import validate_product_xlsx
        from openpyxl import Workbook
        from scholarly import ServiceError
        row=json.loads((ROOT/'tests/fixtures/ams-products.json').read_text('utf-8'))['Products'][0]
        workbook=Workbook();sheet=workbook.active;sheet.append(['Product Code','Title','ISBN','Author/Editor','Publication Date'])
        sheet.append(['OTHER/1','Algebras, Rings and Modules: Lie Algebras and Hopf Algebras','978-0-8218-5262-0',
                      'Michiel Hazewinkel; Nadiya Gubareni; V. V. Kirichenko',2010])
        output=io.BytesIO();workbook.save(output);workbook.close()
        with self.assertRaises(ServiceError) as error:validate_product_xlsx(output.getvalue(),row)
        self.assertEqual(error.exception.result['code'],'identity_mismatch')

    def test_native_xlsx_secret_scan_decodes_xml_entities(self):
        from ams import validate_product_xlsx
        from openpyxl import Workbook
        from scholarly import ServiceError
        row=json.loads((ROOT/'tests/fixtures/ams-products.json').read_text('utf-8'))['Products'][0]
        workbook=Workbook();sheet=workbook.active;sheet.append(['Product Code','Title','ISBN','Author/Editor','Publication Date'])
        sheet.append(['','Algebras, Rings and Modules: Lie Algebras and Hopf Algebras private&credential','978-0-8218-5262-0',
                      'Michiel Hazewinkel; Nadiya Gubareni; V. V. Kirichenko',2010])
        output=io.BytesIO();workbook.save(output);workbook.close()
        with self.assertRaises(ServiceError) as error:validate_product_xlsx(output.getvalue(),row,('private&credential',))
        self.assertEqual(error.exception.result['code'],'sensitive_export')

    def test_selected_window_export_uses_exact_product_search_and_response_body(self):
        from ams import AMS
        from openpyxl import Workbook
        from scholarly import ArtifactStore
        row=json.loads((ROOT/'tests/fixtures/ams-products.json').read_text('utf-8'))['Products'][0]
        workbook=Workbook();sheet=workbook.active;sheet.title='Products'
        sheet.append(['Product Code','Title','ISBN','Author/Editor','Publication Date'])
        for isbn in ('978-0-8218-5262-0','978-1-4704-1395-8',''):
            sheet.append(['','Algebras, Rings and Modules: Lie Algebras and Hopf Algebras',isbn,
                          'Michiel Hazewinkel; Nadiya Gubareni; V. V. Kirichenko',2010])
        output=io.BytesIO();workbook.save(output);workbook.close();xlsx=output.getvalue()
        page=Mock();page.is_closed.return_value=False
        response=Mock(status=200,headers={'content-type':'application/octet-stream','content-disposition':'attachment'},url='https://pubs.ams.org/product/exportproducts')
        response.request=Mock(post_data='advancedSearch=ProductCode%3DSURV%252F168')
        response.body.return_value=xlsx
        pending=Mock();pending.value=response
        page.expect_response.return_value.__enter__=Mock(return_value=pending)
        page.expect_response.return_value.__exit__=Mock(return_value=False)
        page.evaluate.side_effect=lambda script,*args: None
        gateway=Mock();gateway.pages={'ams':page}
        with tempfile.TemporaryDirectory() as temp:
            adapter=AMS(ArtifactStore(temp),gateway)
            def exact(query,**kwargs):
                self.assertEqual((query,kwargs['options'],kwargs['limit'],kwargs['export']),
                                 ('SURV/168',{'field':'ProductCode'},2,None))
                adapter.known['surv/168']=(row,{'ref':'proof'})
                return {'status':'ok','total':1,'received':1,'preview':[row]}
            with patch.object(adapter,'search',side_effect=exact):
                exported=adapter.export_products([row])
            self.assertEqual(exported['status'],'available')
            self.assertEqual(exported['logical_products'],1)
            self.assertEqual(exported['artifacts'][0]['physical_rows'],3)
            self.assertEqual(page.evaluate.call_args.args[0],'uniqueId => initExportResults(uniqueId)')
            self.assertEqual(page.evaluate.call_args.args[1],'')
            self.assertNotIn('headers',exported['artifacts'][0]['artifact'])

    def test_export_can_prepare_institutional_page_after_expression_search(self):
        from ams import AMS
        from openpyxl import Workbook
        from scholarly import ArtifactStore
        row=json.loads((ROOT/'tests/fixtures/ams-products.json').read_text('utf-8'))['Products'][0]
        workbook=Workbook();sheet=workbook.active;sheet.title='Products'
        sheet.append(['Product Code','Title','ISBN','Author/Editor','Publication Date'])
        for isbn in ('978-0-8218-5262-0','978-1-4704-1395-8',''):
            sheet.append(['','Algebras, Rings and Modules: Lie Algebras and Hopf Algebras',isbn,
                          'Michiel Hazewinkel; Nadiya Gubareni; V. V. Kirichenko',2010])
        output=io.BytesIO();workbook.save(output);workbook.close()
        page=Mock();page.is_closed.return_value=False
        response=Mock(status=200,headers={'content-type':'application/octet-stream'},url='https://pubs.ams.org/product/exportproducts')
        response.request=Mock(method='POST',post_data='advancedSearch=ProductCode%3DSURV%252F168')
        response.body.return_value=output.getvalue();pending=Mock();pending.value=response
        page.expect_response.return_value.__enter__=Mock(return_value=pending)
        page.expect_response.return_value.__exit__=Mock(return_value=False)
        gateway=Mock();gateway.pages={}
        with tempfile.TemporaryDirectory() as temp:
            adapter=AMS(ArtifactStore(temp),gateway)
            def exact(query,**kwargs):
                gateway.pages['ams']=page
                return {'status':'ok','total':1,'received':1,'preview':[row]}
            with patch.object(adapter,'search',side_effect=exact) as search:
                exported=adapter.export_products([row])
        self.assertEqual(exported['status'],'available')
        search.assert_called_once_with('SURV/168',options={'field':'ProductCode'},limit=2,export=None)
        self.assertEqual(exported['artifacts'][0]['product_code'],'SURV/168')

    def test_selected_window_export_rejects_stale_product_filter(self):
        from ams import AMS
        from scholarly import ArtifactStore
        row=json.loads((ROOT/'tests/fixtures/ams-products.json').read_text('utf-8'))['Products'][0]
        page=Mock();page.is_closed.return_value=False
        response=Mock(status=200,headers={'content-type':'application/octet-stream'},url='https://pubs.ams.org/product/exportproducts')
        response.request=Mock(post_data='advancedSearch=ProductCode%3DOTHER%252F1')
        pending=Mock();pending.value=response
        page.expect_response.return_value.__enter__=Mock(return_value=pending)
        page.expect_response.return_value.__exit__=Mock(return_value=False)
        gateway=Mock();gateway.pages={'ams':page}
        with tempfile.TemporaryDirectory() as temp:
            adapter=AMS(ArtifactStore(temp),gateway)
            with patch.object(adapter,'search',return_value={'status':'ok','total':1,'received':1,'preview':[row]}):
                exported=adapter.export_products([row])
        self.assertEqual(exported['status'],'error')
        self.assertEqual(exported['errors'][0]['code'],'query_mismatch')

if __name__=='__main__':
    unittest.main()
