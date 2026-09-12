from pathlib import Path
from unittest.mock import MagicMock, patch
import io
import json
import sys
import tempfile
import unittest

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from scholarly import ArtifactStore, ServiceError
from individual_subscription import IndividualSubscription, _issn_status, _url_status


HEADERS = ["TITLE", "ISSN", "E-ISSN", "PUBLISHER", "URL", "COVERAGE\nSTART YEAR", "COVERAGE\nEND YEAR"]


def workbook(rows, *, secret_sources=False):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    if secret_sources:
        marker = "ID" + "/PW: "
        sheet["A2"] = "Example Journal [" + marker + "SYNTHETIC_USER_123 / SYNTHETIC_PASS_456]"
        sheet["D2"].comment = openpyxl.comments.Comment(
            marker + "SYNTHETIC_COMMENT_USER / SYNTHETIC_COMMENT_PASS", "fixture"
        )
        sheet["E2"].hyperlink = "https://other.example/journal?token=SYNTHETIC_LINK_TOKEN"
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def add_comment(data, coordinate, value):
    book = openpyxl.load_workbook(io.BytesIO(data))
    book.active[coordinate].comment = openpyxl.comments.Comment(value, "fixture")
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


class IndividualSubscriptionContract(unittest.TestCase):
    def test_blank_and_styled_rows_do_not_become_holdings_records(self):
        book=openpyxl.Workbook();sheet=book.active;sheet.title='Sheet1'
        sheet.append(HEADERS)
        sheet.append(['One Journal','1234-5679','','Publisher','https://example.org','2000','2026'])
        sheet['A4'].font=openpyxl.styles.Font(bold=True)
        stream=io.BytesIO();book.save(stream);book.close()
        records,metadata=self.adapter._parse_workbook({'catalogue_id':'knu-individual-electronic-journals',
            'url':'https://kudos.knu.ac.kr/download/db/lists/KNU2026.xlsx','data':stream.getvalue()})
        self.assertEqual(len(records),1)
        self.assertEqual(metadata['record_count'],1)
        self.assertEqual(records[0]['row'],2)

    def test_exact_catalogue_fetch_is_independent_of_other_catalogue_failure(self):
        downloads=self.downloads()
        target='knu-individual-electronic-journals'
        def selected_only(selected):
            if selected!=[target]:
                raise ServiceError('synthetic_other_catalogue_down','The other catalogue is unavailable')
            return [downloads[0]]
        with patch.object(self.adapter,'_download_catalogues',side_effect=selected_only):
            result=self.adapter.fetch(target,parts=['metadata'])
        self.assertEqual(result['status'],'ok')
        self.assertEqual(result['catalogue_metadata']['catalogue_id'],target)

    def test_header_only_catalogue_has_no_inverted_data_range(self):
        records,metadata=self.adapter._parse_workbook({'catalogue_id':'knu-individual-electronic-journals',
            'url':'https://kudos.knu.ac.kr/download/db/lists/KNU2026.xlsx','data':workbook([])})
        self.assertEqual(records,[])
        self.assertEqual(metadata['record_count'],0)
        self.assertIsNone(metadata['sheets'][0]['data_coordinate'])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.store = ArtifactStore(self.temp.name)
        self.adapter = IndividualSubscription(self.store, gateway=object())
        self._downloads = None

    def downloads(self):
        if self._downloads is not None:
            return [dict(item) for item in self._downloads]
        current = workbook([
            ["Example", "1234-5679", "1234-5678", "Example Press", " https://bad.example.org /journal ", "All Content", "current"],
            ["SIAM Journal on Discrete Mathematics", "0895-4801", "1095-7146", "SIAM", "https://epubs.siam.org/loi/sjdmec", "1997", "current"],
        ], secret_sources=True)
        archive = workbook([
            ["SIAM Journal on Computing", "0097-5397", "1095-7111", "SIAM", "https://epubs.siam.org/loi/smjcat", "1997", "2025"],
        ])
        self._downloads = [
            {"catalogue_id": "knu-individual-electronic-journals", "url": "https://kudos.knu.ac.kr/download/db/lists/KNU2026.xlsx", "data": current},
            {"catalogue_id": "knu-print-online-journals", "url": "https://kudos.knu.ac.kr/download/db/lists/2025knupo.xlsx", "data": archive},
        ]
        return [dict(item) for item in self._downloads]

    def call(self, method, *args, **kwargs):
        with patch.object(self.adapter, "_download_catalogues", return_value=self.downloads()):
            return getattr(self.adapter, method)(*args, **kwargs)

    def records(self, result):
        return json.loads(self.store.read(result["records"]["ref"])[1])

    def test_credentials_from_cells_hyperlinks_and_comments_never_enter_artifacts(self):
        result = self.call("search", "Example", options={}, limit=10, export="csv")
        self.assertEqual((result["status"], result["received"]), ("ok", 1))
        row = self.records(result)[0]
        self.assertEqual(row["title"], "Example Journal")
        self.assertEqual(row["link_status"], "conflict")
        self.assertEqual(row["displayed_url"], " https://bad.example.org /journal ")
        self.assertEqual(row["displayed_url_status"], "malformed_host")
        self.assertEqual(row["issn_status"], "valid")
        self.assertEqual(row["e_issn"], "1234-5678")
        self.assertEqual(row["e_issn_status"], "invalid_check_digit")
        self.assertEqual(row["hyperlink_target"], "https://other.example/journal")
        self.assertEqual(
            row["redactions"],
            [
                {"coordinate": "Sheet1!A2", "source": "cell"},
                {"coordinate": "Sheet1!D2", "source": "comment"},
                {"coordinate": "Sheet1!E2", "source": "hyperlink"},
            ],
        )
        for path in Path(self.temp.name).rglob("*"):
            if path.is_file():
                data = path.read_bytes()
                for value in (b"SYNTHETIC_USER_123", b"SYNTHETIC_PASS_456", b"SYNTHETIC_COMMENT_USER",
                              b"SYNTHETIC_COMMENT_PASS", b"SYNTHETIC_LINK_TOKEN"):
                    self.assertNotIn(value, data)

    def test_catalogue_source_url_is_sanitized_before_direct_metadata_return(self):
        downloads = self.downloads()
        downloads[1]["url"] += "?token=SYNTHETIC_SOURCE_TOKEN"
        with patch.object(self.adapter, "_download_catalogues", return_value=downloads):
            result = self.adapter.fetch("knu-print-online-journals", parts=["metadata"])
        rendered = json.dumps(result)
        self.assertNotIn("SYNTHETIC_SOURCE_TOKEN", rendered)
        self.assertNotIn("token=", rendered)

    def test_identifiers_and_malformed_urls_are_validated_without_rewriting(self):
        self.assertEqual(_issn_status("BAD1234-5679"), "invalid_format")
        self.assertEqual(_url_status("https://exa mple.org/journal"), "malformed_host")
        self.assertEqual(_url_status("https://[broken.example/journal"), "invalid_url")

    def test_later_catalogue_credentials_redact_earlier_catalogue_fields(self):
        earlier = workbook([
            ["Earlier Journal", "1234-5679", "", "SYNTHETIC_LATER_SECRET", "https://example.org/earlier", "2000", "current"],
        ])
        later = workbook([
            ["Later Journal", "0097-5397", "1095-7111", "Later Press", "https://example.org/later", "2000", "2025"],
        ])
        later = add_comment(later, "A2", "ID" + "/PW: SYNTHETIC_LATER_USER / SYNTHETIC_LATER_SECRET")
        downloads = [
            {"catalogue_id": "knu-individual-electronic-journals", "url": "https://kudos.knu.ac.kr/download/db/lists/first.xlsx", "data": earlier},
            {"catalogue_id": "knu-print-online-journals", "url": "https://kudos.knu.ac.kr/download/db/lists/later.xlsx", "data": later},
        ]
        with patch.object(self.adapter, "_download_catalogues", return_value=downloads):
            result = self.adapter.search("Earlier", options={}, limit=1)
        row = self.records(result)[0]
        self.assertEqual(row["publisher"], "[REDACTED]")
        self.assertIn({"coordinate": "Sheet1!D2", "source": "cell"}, row["redactions"])
        self.assertNotIn("SYNTHETIC_LATER_SECRET", json.dumps(result))

    def test_catalogue_download_uses_redirect_safe_common_downloader(self):
        page = MagicMock()
        page.locator.return_value.count.return_value = 1
        page.locator.return_value.nth.return_value.inner_text.return_value = "국외 KNU 개별전자저널\n제공리스트"
        links = page.locator.return_value.nth.return_value.get_by_role.return_value
        links.count.return_value = 1
        links.first.get_attribute.return_value = "/download/db/lists/KNU2026.xlsx"
        page.url = "https://kudos.knu.ac.kr/pages/sub.htm"
        gateway = MagicMock()
        gateway.context.new_page.return_value = page
        adapter = IndividualSubscription(self.store, gateway)
        data = workbook([["Journal", "1234-5679", "", "Press", "https://example.org", "2000", "current"]])
        with patch("individual_subscription.browser_download", return_value=(data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "https://kudos.knu.ac.kr/download/db/lists/KNU2026.xlsx")) as download:
            result = adapter._download_catalogues(["knu-individual-electronic-journals"])
        self.assertEqual(len(result), 1)
        download.assert_called_once_with(gateway.context, "https://kudos.knu.ac.kr/download/db/lists/KNU2026.xlsx", {"kudos.knu.ac.kr"})
        gateway.context.request.get.assert_not_called()

    def test_literal_search_coverage_filter_exact_range_and_csv_match(self):
        result = self.call("search", "SIAM", options={"field": "all", "coverage_year": 2026}, start=0, limit=1, export="csv")
        self.assertEqual((result["received"], result["total"], result["has_more"]), (1, 1, False))
        row = self.records(result)[0]
        self.assertEqual(row["catalogue_id"], "knu-individual-electronic-journals")
        csv_rows = self.store.read(result["export"]["ref"])[1].decode("utf-8")
        self.assertIn(row["id"], csv_rows)
        self.assertNotIn("SIAM Journal on Computing", csv_rows)

    def test_row_and_catalogue_fetch_are_catalogue_only(self):
        found = self.call("search", "0895-4801", options={"field": "issn"}, limit=2)
        identifier = self.records(found)[0]["id"]
        row = self.call("fetch", identifier, parts=["metadata", "fulltext", "export"])
        self.assertEqual(row["components"]["metadata"]["status"], "available")
        self.assertEqual(row["components"]["export"]["status"], "available")
        self.assertEqual(row["components"]["fulltext"]["status"], "not_offered")
        catalogue = self.call("fetch", "knu-print-online-journals", parts=["metadata", "abstract"])
        self.assertEqual(catalogue["catalogue_metadata"]["record_count"], 1)
        self.assertEqual(catalogue["components"]["abstract"]["status"], "not_offered")
        self.assertNotIn("native_workbook", json.dumps(catalogue))

    def test_invalid_inputs_and_native_workbook_export_are_rejected(self):
        with self.assertRaises(ServiceError) as raised:
            self.call("search", "x", options={"field": "regex"}, limit=1)
        self.assertEqual(raised.exception.result["code"], "invalid_options")
        with self.assertRaises(ServiceError) as raised:
            self.call("search", "x", options={}, limit=1, export="native")
        self.assertEqual(raised.exception.result["code"], "export_not_offered")
        with self.assertRaises(ServiceError) as raised:
            self.call("fetch", "https://epubs.siam.org/loi/sjdmec", parts=["metadata"])
        self.assertEqual(raised.exception.result["code"], "source_boundary")


if __name__ == "__main__":
    unittest.main()
