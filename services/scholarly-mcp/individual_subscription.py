"""KUDOS catalogues for individually subscribed and print-and-online journals."""
from __future__ import annotations
from waits import SOURCE_WAIT_MS

import csv
import io
import json
import re
from datetime import datetime, timezone
from hashlib import sha256
from urllib.parse import parse_qsl, unquote, urljoin, urlsplit

import openpyxl

from browser_adapters import browser_download
from scholarly import ServiceError, encoded


CATALOGUES = {
    "knu-individual-electronic-journals": {
        "label": "국외 KNU 개별전자저널",
        "search_term": "개별전자저널",
        "as_of_year": 2026,
    },
    "knu-print-online-journals": {
        "label": "국외 KNU P+O Journal",
        "search_term": "P+O Journal",
        "as_of_year": 2025,
    },
}
HEADERS = ["TITLE", "ISSN", "E-ISSN", "PUBLISHER", "URL", "COVERAGE\nSTART YEAR", "COVERAGE\nEND YEAR"]
CREDENTIAL_PAIR = re.compile(r"ID\s*/\s*PW\s*:\s*([^\s/\]]+)\s*/\s*([^\s\]]+)", re.I)
CREDENTIAL_ANNOTATION = re.compile(r"\s*\[?\s*ID\s*/\s*PW\s*:\s*[^\s/\]]+\s*/\s*[^\s\]]+\s*\]?", re.I)


def _text(value):
    return "" if value is None else str(value).strip()


def _issn_status(value):
    if not value:
        return "missing"
    match = re.fullmatch(r"(\d{4})-?(\d{3})([0-9Xx])", value)
    if not match:
        return "invalid_format"
    compact = "".join(match.groups()).upper()
    check = (11 - sum(int(compact[index]) * (8 - index) for index in range(7)) % 11) % 11
    expected = "X" if check == 10 else str(check)
    return "valid" if compact[7] == expected else "invalid_check_digit"


def _coverage(value):
    normalized = value.strip()
    if re.fullmatch(r"\d{4}", normalized):
        return {"raw": value, "year": int(normalized), "kind": "year"}
    if normalized.casefold() == "all content":
        return {"raw": value, "year": None, "kind": "all_content"}
    if normalized.casefold() == "current":
        return {"raw": value, "year": None, "kind": "current"}
    return {"raw": value, "year": None, "kind": "invalid"}


def _url_status(value):
    if not value:
        return "missing"
    try:
        parsed = urlsplit(value)
        parsed.port
    except ValueError:
        return "invalid_url"
    if parsed.scheme not in ("http", "https"):
        return "invalid_scheme"
    if (value != value.strip() or not parsed.hostname or any(character.isspace() for character in parsed.netloc)
            or parsed.username or parsed.password):
        return "malformed_host"
    return "valid_https" if parsed.scheme == "https" else "valid_http"


class IndividualSubscription:
    source = "individual-subscription"

    def __init__(self, store, gateway):
        self.store, self.gateway = store, gateway

    def _download_catalogues(self, selected):
        self.gateway.start()
        page = self.gateway.context.new_page()
        results = []
        try:
            self.gateway.login(page)
            page.get_by_role("link", name="검색", exact=True).first.hover()
            page.get_by_role("link", name="학술DB", exact=True).first.click()
            page.locator("#search_order").wait_for(state="visible")
            for catalogue_id in selected:
                details = CATALOGUES[catalogue_id]
                with page.expect_navigation(wait_until="domcontentloaded", timeout=SOURCE_WAIT_MS):
                    page.locator("#search_order").fill(details["search_term"])
                    page.locator("#search_order").press("Enter")
                matches = []
                for index in range(page.locator(".db-tit").count()):
                    row = page.locator(".db-tit").nth(index)
                    lines = [line.strip() for line in row.inner_text().splitlines() if line.strip()]
                    if details["label"] in lines:
                        matches.append(row)
                if len(matches) != 1:
                    raise ServiceError("catalogue_link_not_found", "KUDOS did not expose one exact individual-subscription catalogue row", catalogue=catalogue_id)
                links = matches[0].get_by_role("link", name="제공리스트", exact=True)
                if links.count() != 1:
                    raise ServiceError("catalogue_link_not_found", "KUDOS did not expose one exact catalogue file link", catalogue=catalogue_id)
                target = urljoin(page.url, links.first.get_attribute("href"))
                location = urlsplit(target)
                if (location.scheme != "https" or location.hostname != "kudos.knu.ac.kr" or location.username or location.password
                        or not location.path.startswith("/download/db/lists/") or not location.path.lower().endswith(".xlsx")):
                    raise ServiceError("source_boundary", "KUDOS catalogue link left the validated list-file boundary", catalogue=catalogue_id)
                data, media, actual = browser_download(self.gateway.context, target, {"kudos.knu.ac.kr"})
                actual_location = urlsplit(actual)
                if (actual_location.hostname != "kudos.knu.ac.kr" or not actual_location.path.startswith("/download/db/lists/")
                        or not actual_location.path.lower().endswith(".xlsx")):
                    raise ServiceError("source_boundary", "Catalogue download left the validated KUDOS list-file boundary", catalogue=catalogue_id)
                media = media.split(";", 1)[0].lower()
                if not data.startswith(b"PK") or media != "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
                    raise ServiceError("invalid_workbook", "KUDOS catalogue response is not the expected XLSX workbook", catalogue=catalogue_id)
                results.append({"catalogue_id": catalogue_id, "url": actual, "data": data,
                                "retrieved_at": datetime.now(timezone.utc).isoformat()})
            return results
        finally:
            page.close()

    @staticmethod
    def _workbook_strings(book):
        strings = []
        for sheet in book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        strings.append((sheet.title, cell.coordinate, "cell", str(cell.value)))
                    if cell.hyperlink and cell.hyperlink.target:
                        strings.append((sheet.title, cell.coordinate, "hyperlink", str(cell.hyperlink.target)))
                    if cell.comment:
                        strings.append((sheet.title, cell.coordinate, "comment", str(cell.comment.text or "")))
                        strings.append((sheet.title, cell.coordinate, "comment_author", str(cell.comment.author or "")))
        return strings

    def _register_secrets(self, strings):
        secrets = set()
        for _, _, _, value in strings:
            for match in CREDENTIAL_PAIR.finditer(value):
                secrets.update(part for part in match.groups() if part)
            try:
                parsed = urlsplit(value)
                parsed.port
            except ValueError:
                continue
            if parsed.username:
                secrets.add(unquote(parsed.username))
            if parsed.password:
                secrets.add(unquote(parsed.password))
            for key, item in parse_qsl(parsed.query, keep_blank_values=True):
                if self.store.sanitizer.sensitive.search(key) and item:
                    secrets.add(item)
        self.store.sanitizer.secrets = tuple(sorted(set(self.store.sanitizer.secrets) | secrets, key=len, reverse=True))

    def _prime_workbook_secrets(self, downloaded):
        try:
            book = openpyxl.load_workbook(io.BytesIO(downloaded["data"]), read_only=False, data_only=False)
        except Exception:
            raise ServiceError("invalid_workbook", "KUDOS catalogue workbook could not be parsed", catalogue=downloaded["catalogue_id"]) from None
        try:
            self._register_secrets(self._workbook_strings(book))
        finally:
            book.close()

    def _parse_workbook(self, downloaded):
        data = downloaded["data"]
        try:
            book = openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=False)
        except Exception:
            raise ServiceError("invalid_workbook", "KUDOS catalogue workbook could not be parsed", catalogue=downloaded["catalogue_id"]) from None
        try:
            strings = self._workbook_strings(book)
            self._register_secrets(strings)

            def sanitize(value, *, url=False):
                cleaned = CREDENTIAL_ANNOTATION.sub("", value)
                return self.store.sanitizer.url(cleaned) if url else self.store.sanitizer.text(cleaned)

            redactions = []
            for sheet, coordinate, source, value in strings:
                cleaned = sanitize(value, url=source == "hyperlink")
                if cleaned != value:
                    redactions.append({"coordinate": f"{sheet}!{coordinate}", "source": source})
            redactions = sorted({(item["coordinate"], item["source"]) for item in redactions})
            redactions = [{"coordinate": coordinate, "source": source} for coordinate, source in redactions]

            digest = sha256(data).hexdigest()
            records = []
            sheets = []
            for sheet in book.worksheets:
                headers = [_text(cell.value) for cell in sheet[1]]
                if headers != HEADERS:
                    raise ServiceError("invalid_workbook_schema", "KUDOS catalogue workbook headers changed", catalogue=downloaded["catalogue_id"], sheet=sheet.title)
                sheets.append({"name": sheet.title, "header_coordinate": f"{sheet.title}!A1:G1",
                               "data_coordinate": f"{sheet.title}!A2:G{sheet.max_row}" if sheet.max_row>=2 else None,
                               "coordinate_scope": "Physical worksheet extent; blank rows are not holdings records"})
                for row_number in range(2, sheet.max_row + 1):
                    cells = [sheet.cell(row_number, column) for column in range(1, 8)]
                    if not any(_text(cell.value) or cell.hyperlink and cell.hyperlink.target for cell in cells):
                        continue
                    values = [sanitize("" if cell.value is None else str(cell.value)) for cell in cells]
                    title, issn, e_issn, publisher, displayed_url, start_raw, end_raw = values
                    hyperlink = cells[4].hyperlink.target if cells[4].hyperlink and cells[4].hyperlink.target else None
                    hyperlink = sanitize(str(hyperlink), url=True) if hyperlink else None
                    start = _coverage(start_raw)
                    end = _coverage(end_raw)
                    row_redactions = [item for item in redactions if item["coordinate"].startswith(f"{sheet.title}!")
                                      and re.search(rf"[A-Z]+{row_number}$", item["coordinate"])]
                    link_status = "displayed_only" if hyperlink is None else "matching" if hyperlink == displayed_url else "conflict"
                    records.append({
                        "id": f"{downloaded['catalogue_id']}:{digest[:12]}:{sheet.title}:{row_number}",
                        "catalogue_id": downloaded["catalogue_id"], "workbook_sha256": digest,
                        "sheet": sheet.title, "row": row_number, "title": title, "issn": issn, "e_issn": e_issn,
                        "issn_status": _issn_status(issn), "e_issn_status": _issn_status(e_issn),
                        "publisher": publisher, "displayed_url": displayed_url,
                        "displayed_url_status": _url_status(displayed_url), "hyperlink_target": hyperlink,
                        "link_status": link_status, "coverage_start_raw": start["raw"],
                        "coverage_start_year": start["year"], "coverage_start_kind": start["kind"],
                        "coverage_end_raw": end["raw"], "coverage_end_year": end["year"],
                        "coverage_end_kind": end["kind"], "redactions": row_redactions,
                    })
            metadata = {
                "catalogue_id": downloaded["catalogue_id"], "label": CATALOGUES[downloaded["catalogue_id"]]["label"],
                "source_url": self.store.sanitizer.url(downloaded["url"]), "retrieved_at": downloaded.get("retrieved_at"),
                "workbook_sha256": digest, "workbook_bytes": len(data), "record_count": len(records),
                "sheets": sheets, "columns": HEADERS, "redactions": redactions,
                "raw_workbook_exposed": False,
            }
            return records, metadata
        finally:
            book.close()

    def _load(self, selected=None):
        selected = list(CATALOGUES) if selected is None else list(selected)
        records, metadata = [], []
        downloaded = self._download_catalogues(selected)
        for item in downloaded:
            if item["catalogue_id"] in selected:
                self._prime_workbook_secrets(item)
        for item in downloaded:
            if item["catalogue_id"] not in selected:
                continue
            parsed, description = self._parse_workbook(item)
            records.extend(parsed)
            metadata.append(description)
        if {item["catalogue_id"] for item in metadata} != set(selected):
            raise ServiceError("catalogue_missing", "KUDOS did not return every requested catalogue")
        return self.store.sanitizer.value(records), self.store.sanitizer.value(metadata)

    @staticmethod
    def _covers(record, year):
        if year is None:
            return True
        if record["coverage_start_kind"] == "invalid" or record["coverage_end_kind"] == "invalid":
            return False
        if record["coverage_start_kind"] == "year" and year < record["coverage_start_year"]:
            return False
        if record["coverage_end_kind"] == "year" and year > record["coverage_end_year"]:
            return False
        if record["coverage_end_kind"] == "current" and year > CATALOGUES[record["catalogue_id"]]["as_of_year"]:
            return False
        return True

    def search(self, query, *, options, start=0, limit=20, cursor=None, export=None):
        allowed = {"catalogue", "field", "coverage_year"}
        if set(options) - allowed or cursor:
            raise ServiceError("invalid_options", "Use catalogue, field and coverage_year; pagination uses start")
        if not isinstance(query, str):
            raise ServiceError("invalid_query", "Catalogue query must be a string")
        if type(start) is not int or start < 0 or type(limit) is not int or not 1 <= limit <= 1000:
            raise ServiceError("invalid_range", "Use start >= 0 and limit 1..1000")
        if export == "native":
            raise ServiceError("export_not_offered", "Raw catalogue workbooks are not exposed because they can contain credentials")
        if export not in (None, "json", "csv"):
            raise ServiceError("invalid_export", "Use json or csv export")
        catalogue = options.get("catalogue", "all")
        if catalogue != "all" and catalogue not in CATALOGUES:
            raise ServiceError("invalid_options", "Use all or a documented catalogue identifier", field="catalogue")
        field = options.get("field", "all")
        if field not in ("all", "title", "publisher", "issn", "url"):
            raise ServiceError("invalid_options", "Use all, title, publisher, issn or url", field="field")
        year = options.get("coverage_year")
        if year is not None and (type(year) is not int or not 1000 <= year <= 9999):
            raise ServiceError("invalid_options", "coverage_year must be a four-digit integer", field="coverage_year")
        selected_catalogues = list(CATALOGUES) if catalogue == "all" else [catalogue]
        rows, descriptions = self._load(selected_catalogues)
        needle = query.strip().casefold()
        def haystack(row):
            fields = {
                "title": row["title"], "publisher": row["publisher"],
                "issn": " ".join((row["issn"], row["e_issn"])),
                "url": " ".join(value for value in (row["displayed_url"], row["hyperlink_target"]) if value),
            }
            return fields[field] if field != "all" else " ".join(fields.values())
        matched = [row for row in rows if (not needle or needle in haystack(row).casefold()) and self._covers(row, year)]
        window = matched[start:start + limit]
        source_url = descriptions[0]["source_url"] if len(descriptions) == 1 else "https://kudos.knu.ac.kr/"
        catalogue_metadata = self.store.put(encoded(descriptions), "application/json", self.source, source_url,
                                            role="catalogue_metadata", catalogues=selected_catalogues)
        records = self.store.put(encoded(window), "application/json", self.source, source_url, role="records",
                                 query=query, options=options, start=start, requested=limit,
                                 total=len(matched), catalogue_metadata_ref=catalogue_metadata["ref"])
        exported = records
        if export == "csv":
            stream = io.StringIO(newline="")
            fieldnames = list(window[0]) if window else ["id", "catalogue_id", "title", "issn", "e_issn"]
            writer = csv.DictWriter(stream, fieldnames=fieldnames)
            writer.writeheader()
            for row in window:
                writer.writerow({key: json.dumps(value, ensure_ascii=False, separators=(",", ":")) if isinstance(value, (list, dict))
                                 else "" if value is None else value for key, value in row.items()})
            exported = self.store.put(stream.getvalue(), "text/csv", self.source, source_url,
                                      role="catalogue_csv", source_format="mechanical_csv", parent_ref=records["ref"])
        more = start + len(window) < len(matched)
        return {"status": "ok" if window else "not_found", "source": self.source, "query": query, "options": options,
                "start": start, "requested": limit, "received": len(window), "total": len(matched), "has_more": more,
                "next": {"start": start + len(window)} if more else None, "records": records, "export": exported,
                "catalogue_metadata": catalogue_metadata,
                "preview": [{key: row[key] for key in ("id", "title", "issn", "e_issn", "catalogue_id")} for row in window[:3]],
                "scope": "KUDOS journal holdings catalogue; external publisher URLs are unvisited candidates"}

    def fetch(self, target, *, parts):
        if set(parts) - {"metadata", "abstract", "fulltext", "supplementary", "export"}:
            raise ServiceError("invalid_parts", "Use documented content parts")
        if target.startswith(("http://", "https://")):
            raise ServiceError("source_boundary", "Use a catalogue or row identifier; publisher URLs are separate candidate sources")
        catalogue_id=target.split(':',1)[0]
        if catalogue_id not in CATALOGUES:
            return {"status":"not_found","source":self.source,"target":target,"message":"Unknown catalogue or row identifier"}
        rows, descriptions = self._load([catalogue_id])
        description = next((item for item in descriptions if item["catalogue_id"] == target), None)
        row = next((item for item in rows if item["id"] == target), None)
        if description is None and row is None:
            return {"status": "not_found", "source": self.source, "target": target, "message": "Catalogue record was not found in the current workbooks"}
        value = description if description is not None else row
        artifact = self.store.put(encoded([value]), "application/json", self.source,
                                  description["source_url"] if description else "https://kudos.knu.ac.kr/",
                                  role="catalogue_metadata" if description else "records", target=target)
        components = {}
        for part in parts:
            if part in ("metadata", "export"):
                components[part] = {"status": "available", "artifact": artifact,
                                    "format": "sanitized_json" if part == "export" else None}
            else:
                components[part] = {"status": "not_offered",
                                    "reason": "This source is a holdings catalogue; select a publisher source explicitly for article content"}
        status = "ok" if all(item["status"] == "available" for item in components.values()) else "partial"
        result = {"status": status, "source": self.source, "target": target, "records": artifact, "components": components,
                  "scope": "Sanitized KUDOS holdings metadata; no publisher URL was visited"}
        if description is not None:
            result["catalogue_metadata"] = description
        return result
