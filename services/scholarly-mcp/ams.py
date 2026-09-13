"""AMS Digital Library native institutional browse records."""
import csv,html,io,json,re,unicodedata,zipfile
from urllib.parse import urlsplit,urljoin,parse_qs,urlencode,urlunsplit,unquote
from browser_adapters import goto_source,clean_html,dismiss_cookies,is_access_challenge
from publishers import Publisher,PROFILES
from scholarly import ServiceError,encoded


def advanced_expression(expression,_nested=False):
    """Translate only documented controls; preserve groups, values and operators."""
    if not isinstance(expression,dict):raise ServiceError('invalid_options','An AMS expression must be an object')
    if set(expression)=={'condition','terms'}:
        if expression['condition'] not in ('AND','OR') or not isinstance(expression['terms'],list) or not expression['terms']:
            raise ServiceError('invalid_options','Expression groups need AND or OR and a nonempty terms list')
        return {'condition':expression['condition'],'queryElements':[advanced_expression(term,True) for term in expression['terms']],
                **({'field':'','operator':'','value':''} if _nested else {})}
    fields={'Title':'ProductTitle','Author':'Author','ProductCode':'ProductCode','BookText':'Anywhere',
            'PublicationYear':'PublicationYear','ISBN':'ISBN','MSC':'MSC','Series':'BookSeries','Subject':'Subject'}
    if set(expression)!={'field','operator','value'} or not isinstance(expression['field'],str) or expression['field'] not in fields:
        raise ServiceError('invalid_options','Use an AMS expression group or a documented field/operator/value term')
    operators=('equals','contains') if expression['field'] in ('Title','Author','ProductCode','BookText') else ('equals',)
    if expression['operator'] not in operators or not isinstance(expression['value'],str) or not expression['value'].strip():
        raise ServiceError('invalid_options','The term needs a supported field operator and a nonempty string value')
    return {'condition':'','queryElements':None,**expression,'field':fields[expression['field']]}


def basic_expression(expression):
    """Compile the provider-verified expression subset to visible basic controls."""
    advanced_expression(expression)

    def unavailable():
        raise ServiceError(
            'provider_feature_unavailable',
            'The current AMS controls cannot execute this Boolean field combination',
            status='unsupported',source='ams',stage='search_options',
            supported_expression_subset={
                'terms':'Title or Author with contains and one alphanumeric token',
                'groups':'Nested same-field OR, or AND over one Title expression and one Author expression',
            })

    def unmapped():
        raise ServiceError(
            'expression_mapping_unavailable',
            'The local AMS adapter has not mapped this field, operator, or atomic value without changing its meaning',
            status='unsupported',source='ams',stage='search_options')

    def term(value):
        if set(value)!={'field','operator','value'} or value['field'] not in ('Title','Author') or value['operator']!='contains':
            unmapped()
        text=unicodedata.normalize('NFC',value['value'].strip())
        if not text.isalnum() or text.upper() in ('AND','OR','NOT'):
            unmapped()
        return value['field'],text

    def compile_node(value):
        if set(value)=={'field','operator','value'}:
            field,text=term(value)
            return {field:text}
        compiled=[compile_node(child) for child in value['terms']]
        if value['condition']=='OR':
            fields={next(iter(item)) for item in compiled if len(item)==1}
            if len(fields)!=1 or len(compiled)!=sum(len(item)==1 for item in compiled):
                unavailable()
            field=fields.pop()
            return {field:' '.join(item[field] for item in compiled)}
        result={}
        for item in compiled:
            if set(result)&set(item):
                unavailable()
            result.update(item)
        if not result or set(result)-{'Title','Author'}:
            unavailable()
        return result

    return compile_node(expression)


def basic_search_value(value):
    """Mirror the visible AMS form's documented client-side value normalization."""
    decomposed=unicodedata.normalize('NFD',value)
    return ''.join(char for char in decomposed if not '\u0300'<=char<='\u036f').replace('\u2018',"'").replace('\u2019',"'")


def basic_search_expected(values,fields):
    return {fields[name][1]:[basic_search_value(value)] for name,value in values.items()}


def check_expression_request(body,expected):
    request=parse_qs(body or '');values=request.get('advancedSearchJson',[])+request.get('dynamicAdvancedSearch',[])
    if len(values)==1 and not request.get('advancedSearch'):
        value=values[0]
        for _ in range(3):
            try:observed=json.loads(value)
            except ValueError:value=unquote(value);continue
            if observed==expected:return
            break
    raise ServiceError('query_mismatch','AMS did not retain the submitted expression and groups')


def basic_search_request(body):
    """Read the provider product API's pipe-delimited advanced-search fields."""
    request=parse_qs(body or '')
    values=request.get('advancedSearch',[])
    return parse_qs(values[0],separator='|') if len(values)==1 else {}


def prepare_expression(page,expression):
    page.wait_for_function("() => typeof showDynamic==='function' && typeof loadDynamicSearch==='function' && typeof getQueryElements==='function'",timeout=20000)
    observed=page.evaluate('''expression => {
        showDynamic(); loadDynamicSearch(JSON.stringify(expression));
        return {condition: document.querySelector('.advancedSearchGroupHeader[data-group-id="1"][data-group-parent-id="0"] button.selected')?.dataset.conditionValue,
                queryElements: getQueryElements(0,1)};
    }''',expression)
    if observed!=expression:
        raise ServiceError('query_not_supported','The current AMS controls cannot represent the requested field values or groups',status='unsupported')


def validate_search_page(raw,http_status=None):
    from bs4 import BeautifulSoup
    soup=BeautifulSoup(raw,'html.parser')
    title=soup.title.get_text(' ',strip=True) if soup.title else ''
    body=soup.get_text(' ',strip=True)
    if http_status in (401,403) or is_access_challenge(raw):
        raise ServiceError('publisher_access_challenge','AMS did not provide an accessible search page',status='restricted',source='ams',stage='search_interface',http_status=http_status)
    if body.strip()=='The service is unavailable.' or http_status is not None and http_status>=500:
        raise ServiceError('provider_unavailable','AMS returned a service-unavailable page',source='ams',stage='search_interface',http_status=http_status,retryable=True)
    if http_status is not None and http_status>=400:
        raise ServiceError('search_interface_unavailable','AMS did not provide its search interface',source='ams',stage='search_interface',http_status=http_status)


def product_records(payload):
    if not isinstance(payload,dict):
        raise ServiceError('invalid_response','AMS returned an error or non-object response instead of product results')
    rows=payload.get('Products')
    total,page,capacity=(payload.get(key) for key in ('TotalRows','PageNumber','RowsPerPage'))
    if rows is None and type(total) is int and total==0:
        rows=[]
    if not isinstance(rows,list) or any(not isinstance(row,dict) for row in rows) or any(type(v) is not int for v in (total,page,capacity)):
        raise ServiceError('invalid_response','AMS returned no native product range')
    normalized=[]
    for row in rows:
        if row.get('Title'):
            normalized.append(row)
            continue
        reviews=row.get('BookReviews')
        if row.get('DocumentType')=='BOOKREV' and isinstance(reviews,list) and len(reviews)==1 and isinstance(reviews[0],dict) and reviews[0].get('Title'):
            normalized.append({**row,'Title':reviews[0]['Title'],'TitleSource':'BookReviews[0].Title'})
            continue
        raise ServiceError('invalid_response','AMS returned a product without a native display title')
    remaining=max(0,total-(page-1)*capacity)
    if (total<0 or page<1 or capacity<1 or len(rows)>min(capacity,remaining)
            or bool(rows)!=(remaining>0)):
        raise ServiceError('range_mismatch','AMS products do not match the declared range')
    return normalized,total,(page-1)*capacity


def requested_page_gap(first,count,capacity,total,start,limit):
    """Describe a provider-declared ordinal hole only when it intersects the request."""
    actual_end=first+count
    declared_end=min(first+capacity,total)
    if actual_end>=declared_end or actual_end>=start+limit or declared_end<=start:
        return None
    return {'status':'error','code':'source_range_gap',
            'reason':'AMS displayed fewer products than its declared page range',
            'missing_start':max(start,actual_end),'missing_end':min(start+limit,declared_end),
            'resume_start':declared_end,'page_first':first,'page_received':count,
            'page_capacity':capacity,'total':total}


def native_record_fingerprint(row):
    """Use the whole native row: AMS ProductId can identify several abstract rows."""
    return encoded(row)


def validate_page_sequence(observed_page,expected_page,previous_total,reported_total):
    if observed_page!=expected_page or previous_total is not None and previous_total!=reported_total:
        raise ServiceError('range_mismatch','AMS changed the page or result total')


def select_native_page(rows,first,capacity,total,start,limit,seen):
    """Select a confirmed ordinal prefix and report the first source-side hole."""
    selected=[]
    for index,row in enumerate(rows,first):
        fingerprint=native_record_fingerprint(row)
        duplicate=fingerprint in seen
        seen.add(fingerprint)
        if start<=index<start+limit:
            if duplicate:
                return selected,{'status':'error','code':'source_duplicate_record',
                    'reason':'AMS repeated an identical native product at another ordinal',
                    'duplicate_at':index,'resume_start':index+1,'total':total}
            selected.append((index,row))
    return selected,requested_page_gap(first,len(rows),capacity,total,start,limit)


def page_export_scope(format,incomplete):
    if incomplete:
        return 'Visited native pages; records contains only the confirmed prefix; the requested window is incomplete'
    if format=='native':
        return 'Whole native pages covering the requested window'
    return 'Visited native pages; records contains the exact requested window'


def validate_product_xlsx(data,row,secrets=()):
    """Validate native variant rows without filling absent provider values."""
    from defusedxml import ElementTree
    from openpyxl import load_workbook
    if not isinstance(data,bytes) or not data.startswith(b'PK\x03\x04'):
        raise ServiceError('invalid_native_export','AMS did not return an XLSX file')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names=archive.namelist()
            if not all(name in names for name in ('[Content_Types].xml','xl/workbook.xml')):
                raise ServiceError('invalid_native_export','AMS export is not an OOXML workbook')
            protected=[]
            for name in names:
                if not name.endswith(('.xml','.rels')):continue
                raw=archive.read(name)
                decoded=html.unescape(raw.decode('utf-8',errors='replace'))
                protected.append(decoded)
                root=ElementTree.fromstring(raw)
                for element in root.iter():
                    protected.extend(value for value in (element.text,element.tail,*element.attrib.keys(),*element.attrib.values()) if value)
        inspected='\n'.join(protected).casefold()
        known=[str(value).casefold() for value in secrets if isinstance(value,str) and value]
        markers=('__requestverificationtoken','name="password"','name="passwd"','name="cookie"')
        if any(value in inspected for value in known) or any(marker in inspected for marker in markers):
            raise ServiceError('sensitive_export','AMS XLSX contains a protected credential or request marker')
        workbook=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        sheets=[]
        for sheet in workbook.worksheets:
            rows=list(sheet.iter_rows(values_only=True))
            if rows:sheets.append((sheet.title,rows))
        workbook.close()
    except ServiceError:raise
    except Exception:
        raise ServiceError('invalid_native_export','AMS export is not a readable XLSX workbook') from None
    if len(sheets)!=1:
        raise ServiceError('invalid_native_export','AMS native product export must contain one populated sheet')
    sheet_name,rows=sheets[0]
    expected_headers=('Product Code','Title','ISBN','Author/Editor','Publication Date')
    headers=tuple(str(value) if value is not None else '' for value in rows[0])
    if headers!=expected_headers or len(rows)<2:
        raise ServiceError('invalid_native_export','AMS native product export has an unknown table structure')
    physical=[[str(value) if value is not None else '' for value in values] for values in rows[1:]]
    title=str(row.get('Title') or '').strip();subtitle=str(row.get('Subtitle') or '').strip()
    expected_title=title+(': '+subtitle if subtitle else '')
    if not expected_title or any(values[1]!=expected_title for values in physical):
        raise ServiceError('identity_mismatch','AMS XLSX titles do not match the selected product')
    groups={(values[1],values[3],values[4]) for values in physical}
    if len(groups)!=1:
        raise ServiceError('identity_mismatch','AMS XLSX rows do not identify one logical product')
    native_isbn=str(row.get('ISBN') or '').strip()
    isbns=[values[2] for values in physical]
    if native_isbn and native_isbn not in isbns:
        raise ServiceError('identity_mismatch','AMS XLSX does not contain the selected product ISBN')
    product_code=str(row.get('ProductCode') or '').strip()
    if not product_code or any(values[0].strip() not in ('',product_code) for values in physical):
        raise ServiceError('identity_mismatch','AMS XLSX Product Code cells conflict with the selected product')
    related=row.get('RelatedCustomers') or []
    expected_people=[html.unescape(str(person.get('FullName') or '')).strip() for person in related
                     if isinstance(person,dict) and person.get('RelationshipType') in ('AUTHOR','EDITOR') and person.get('FullName')]
    exported_people=[part.strip() for part in physical[0][3].split(';') if part.strip()]
    if expected_people and exported_people!=expected_people:
        raise ServiceError('identity_mismatch','AMS XLSX authors or editors do not match the selected product')
    expected_date=row.get('CopyrightYear')
    if expected_date is not None and any(values[4]!=str(expected_date) for values in physical):
        raise ServiceError('identity_mismatch','AMS XLSX publication date does not match the selected product')
    return {'verified':True,'format':'xlsx','sheet':sheet_name,'headers':list(headers),'physical_rows':len(physical),
            'logical_products':1,'product_code':product_code,'product_code_cells':sum(bool(values[0]) for values in physical),
            'native_isbn_rows':sum(values[2]==native_isbn for values in physical) if native_isbn else None,
            'variant_isbns':isbns,'scope':'One selected logical product; provider ISBN variants remain separate physical rows'}


class AMS(Publisher):
    fields={'Title':('advancedSearchTitleInput','ProductTitle'),'Author':('advancedSearchAuthorInput','Author'),
            'ProductCode':('advancedSearchFreeformInput','ProductCode'),'BookText':('advancedSearchAnywhereInput','Anywhere'),
            'ISBN':('advancedSearchISBNInput','ISBN'),'MSC':('advancedSearchMSCInput','MSC')}

    def __init__(self,store,gateway):
        super().__init__(store,gateway,next(p for p in PROFILES if p.source=='ams'))
        self.known={}

    def export_products(self,selected):
        """Export the exact selected window as one native workbook per product."""
        artifacts=[];errors=[]
        for row in selected:
            code=row.get('ProductCode')
            if not isinstance(code,str) or not code.strip():
                errors.append({'status':'unsupported','code':'native_export_identity_unavailable','reason':'AMS native export needs a ProductCode that can select one logical product'})
                continue
            try:
                exact=self.search(code,options={'field':'ProductCode'},limit=2,export=None)
                if exact.get('total')!=1 or exact.get('received')!=1:
                    raise ServiceError('native_export_identity_ambiguous','AMS ProductCode does not select exactly one logical product',status='unsupported')
                exact_row=(exact.get('preview') or [{}])[0]
                if exact_row.get('ProductCode')!=code:
                    raise ServiceError('identity_mismatch','AMS exact export search returned another product')
                page=self.gateway.pages.get(self.source)
                if page is None or page.is_closed():
                    raise ServiceError('export_page_unavailable','AMS did not prepare the institutional product page for export')
                def matches(response):
                    parts=urlsplit(response.url)
                    return parts.hostname in self.profile.hosts and parts.path=='/product/exportproducts' and response.request.method=='POST'
                with page.expect_response(matches,timeout=60000) as pending:
                    page.evaluate('uniqueId => initExportResults(uniqueId)','')
                response=pending.value
                if response.status in (401,403):
                    raise ServiceError('access_restricted','AMS restricted the native product export',status='restricted')
                if response.status!=200:
                    raise ServiceError('provider_error','AMS native product export failed',http_status=response.status)
                request=parse_qs(response.request.post_data or '')
                advanced=parse_qs(request.get('advancedSearch',[''])[0])
                if advanced.get('ProductCode')!=[code]:
                    raise ServiceError('query_mismatch','AMS native export did not retain the exact ProductCode selection')
                data=response.body()
                integrity=validate_product_xlsx(data,row,getattr(self.store.sanitizer,'secrets',()))
                actual=urlsplit(response.url);safe_url=urlunsplit((actual.scheme,actual.netloc,actual.path,'',''))
                artifact=self.store.put(data,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',self.source,safe_url,
                    preserve_bytes=True,role='native_product_export',requested_product_code=code,integrity=integrity)
                artifacts.append({'product_code':code,'artifact':artifact,**integrity})
            except ServiceError as error:
                errors.append({'product_code':code,**error.result})
            except Exception as error:
                errors.append({'product_code':code,'status':'error','code':'native_export_failed','error_type':type(error).__name__})
        status='available' if artifacts and not errors else 'partial' if artifacts else 'error'
        return {'status':status,'artifacts':artifacts,'errors':errors,'logical_products':len(selected),'available_products':len(artifacts),
                'scope':'One provider XLSX per selected logical product; physical rows retain native ISBN variants'}

    def search(self,query,*,options,start=0,limit=20,cursor=None,export=None):
        if type(start) is not int or start<0 or type(limit) is not int or not 1<=limit<=1000:
            raise ServiceError('invalid_range','Use start >= 0 and limit 1..1000')
        if cursor:
            raise ServiceError('invalid_options','AMS search does not accept a cursor; use start and limit')
        expression_values=None
        if 'expression' in options:
            if query or set(options)!={'expression'}:
                raise ServiceError('invalid_options','Expression cannot be combined with a query or other AMS options')
            expression_values=basic_expression(options['expression'])
        if set(options)-{'field','title','author','product_code','book_text','isbn','msc','content_type'} or options.get('field','Title') not in self.fields:
            if expression_values is None:
                raise ServiceError('invalid_options','Use the documented AMS advanced-search fields')
        expression=None
        if export not in (None,'json','native','csv','html','xlsx'):
            raise ServiceError('invalid_export','Use json, native, csv, html or the verified native xlsx export')
        content=options.get('content_type','All')
        if content not in ('All','EBook','Journal','Abstracts'):
            raise ServiceError('invalid_options','content_type must be All, EBook, Journal or Abstracts')
        if expression_values is not None:
            public_key=self.source+'-expression-public'
            cached=self.gateway.pages.get(public_key)
            if cached is not None and not cached.is_closed():
                cached.close()
                self.gateway.pages.pop(public_key,None)
            page=self.gateway.public_page(public_key)
            search_page='https://pubs.ams.org/AdvancedSearch'
        else:
            page=self.gateway.open('ams','AMS',self.profile.hosts,'body')
            search_page='https://'+urlsplit(page.url).hostname+'/AdvancedSearch'
        try:
            response=goto_source(page,search_page,self.profile.hosts,wait_until='domcontentloaded',timeout=45000)
        except Exception as first_error:
            from playwright.sync_api import TimeoutError as BrowserTimeout
            if not isinstance(first_error,BrowserTimeout):raise
            current=urlsplit(page.url)
            ready=(current.hostname in self.profile.hosts and current.path=='/AdvancedSearch'
                   and page.locator('#advancedSearchTitleInput').count())
            if ready:
                response=None
            else:
                try:
                    response=goto_source(page,search_page,self.profile.hosts,wait_until='commit',timeout=20000)
                except Exception as retry_error:
                    raise ServiceError('search_interface_not_ready','AMS advanced-search navigation did not stabilize',
                                       source=self.source,stage='search_interface_navigation',
                                       initial_error_type=type(first_error).__name__,
                                       retry_error_type=type(retry_error).__name__) from None
        if response is not None:
            validate_search_page(page.content(),response.status)
        try:
            page.wait_for_function('(ids) => ids.every(id=>document.getElementById(id))',arg=[control for control,_ in self.fields.values()],timeout=20000)
        except Exception as error:
            validate_search_page(page.content())
            raise ServiceError('search_interface_not_ready','AMS did not initialize the documented advanced-search controls',source=self.source,stage='search_interface',error_type=type(error).__name__) from None
        if response is None:
            validate_search_page(page.content())
        okay=page.get_by_role('button',name='Okay',exact=True).filter(visible=True)
        if okay.count():
            okay.first.evaluate('node => node.click()')
        else:
            try:
                dismiss_cookies(page)
            except Exception as error:
                okay=page.get_by_role('button',name='Okay',exact=True).filter(visible=True)
                if okay.count():
                    okay.first.evaluate('node => node.click()')
                else:
                    raise ServiceError('search_interface_not_ready','AMS cookie controls did not become actionable',
                                       source=self.source,stage='cookie_control',error_type=type(error).__name__) from None
        for control,_ in self.fields.values():
            if page.locator('#'+control).input_value():
                page.locator('#'+control).fill('')
        field=options.get('field','Title')
        values=expression_values if expression_values is not None else {field:query}
        aliases={'title':'Title','author':'Author','product_code':'ProductCode','book_text':'BookText','isbn':'ISBN','msc':'MSC'}
        for key,name in aliases.items():
            if key in options:
                if name in values:
                    raise ServiceError('invalid_options','The primary field and an additional field must not overlap',field=name)
                values[name]=str(options[key])
        if expression:
            prepare_expression(page,expression)
        else:
            for name,value in values.items():
                page.locator('#'+self.fields[name][0]).fill(value)
            if not page.locator('#advancedSearchContentTypeInput'+content).is_checked():
                page.locator('label[for="advancedSearchContentTypeInput'+content+'"]').evaluate('node => node.click()')

        def matches(response):
            return urlsplit(response.url).hostname in self.profile.hosts and urlsplit(response.url).path=='/product/browseproducts' and response.request.method=='POST'

        try:
            with page.expect_response(matches,timeout=60000) as pending:
                page.locator('#advancedSearchGo').filter(visible=True).filter(has_text=re.compile(r'^Search$')).first.evaluate('node => node.click()')
                if re.search(r'^(Just a moment|Access Denied|Attention Required)',page.title(),re.I):
                    raise ServiceError('publisher_access_challenge','AMS displayed an access verification page',status='restricted')
        except ServiceError:
            raise
        except Exception as error:
            validate_search_page(page.content())
            challenged=bool(re.search(r'^(Just a moment|Access Denied|Attention Required)',page.title(),re.I))
            raise ServiceError('publisher_access_challenge' if challenged else 'search_response_not_received','AMS did not provide a native search response',status='restricted' if challenged else 'error',error_type=type(error).__name__) from None
        response=pending.value
        selected,evidence,html_pages=[],[],[]
        errors=[];seen=set();incomplete=False;resume_start=None;last_capacity=None
        total=None
        search_url=None
        page_number=1
        while True:
            try:
                if response.status in (401,403):
                    raise ServiceError('access_restricted','AMS restricted this search response',status='restricted')
                if response.status!=200:
                    raise ServiceError('provider_error','AMS search failed',http_status=response.status)
                advanced=basic_search_request(response.request.post_data)
                if expression:
                    check_expression_request(response.request.post_data,expression)
                elif any(advanced.get(field)!=value for field,value in basic_search_expected(values,self.fields).items()):
                    raise ServiceError('query_mismatch','AMS native request does not retain the submitted fields')
                payload=response.json()
                native,reported,first=product_records(payload)
                expected_type={'EBook':'B','Journal':'A','Abstracts':'ABS'}.get(content)
                if expected_type and any(row.get('ContentType')!=expected_type for row in native):
                    raise ServiceError('filter_mismatch','AMS returned a product outside the requested content type')
                validate_page_sequence(payload['PageNumber'],page_number,total,reported)
            except ServiceError as error:
                if not evidence:raise
                try:
                    failed=self.store.put(response.body(),response.headers.get('content-type','application/octet-stream'),self.source,response.url,
                                          role='native_search_error_response',query=query,options=options,page_number=page_number)
                    evidence.append(failed)
                except Exception:pass
                resume_start=(page_number-1)*(last_capacity or 1)
                errors.append({**error.result,'page_number':page_number,'resume_start':resume_start})
                incomplete=True
                break
            total=reported
            last_capacity=payload['RowsPerPage']
            proof=self.store.put(encoded(payload),'application/json',self.source,response.url,role='native_search_response',query=query,options=options)
            evidence.append(proof)
            selector='#productBrowseListContainer > li.productBrowseListItem:not(.placeHolderItem) a.productBrowseTitleLink'
            try:
                try:
                    page.wait_for_function('(arg) => document.querySelectorAll(arg.selector).length === arg.count',arg={'selector':selector,'count':len(native)},timeout=30000)
                except Exception as error:
                    raise ServiceError('search_render_mismatch','AMS rendered a different number of product links',
                                       page_number=page_number,native_count=len(native),
                                       rendered_count=page.locator(selector).count(),error_type=type(error).__name__) from None
                if native:
                    expected='Showing '+str(first+1)+' - '+str(first+len(native))+' of '+str(total)+' results'
                    page.wait_for_function('(expected) => document.querySelector("#productBrowseResultCountTopOneColumn")?.textContent.replace(/\\s+/g," ").trim() === expected',arg=expected,timeout=30000)
            except ServiceError as error:
                if page_number==1:raise
                resume_start=first
                errors.append({**error.result,'page_number':page_number,'resume_start':resume_start})
                incomplete=True
                break
            except Exception as error:
                if native:
                    count=page.locator('#productBrowseResultCountTopOneColumn')
                    observed=' '.join(count.inner_text().split()) if count.count() else None
                    failure=ServiceError('search_render_mismatch','AMS rendered a different product ordinal range',
                                         page_number=page_number,expected=expected,observed=observed,
                                         error_type=type(error).__name__)
                    if page_number==1:raise failure from None
                    resume_start=first
                    errors.append({**failure.result,'page_number':page_number,'resume_start':resume_start})
                    incomplete=True
                    break
            page_selected,issue=select_native_page(native,first,payload['RowsPerPage'],total,start,limit,seen)
            selected_indices={index for index,_ in page_selected}
            anchors=page.locator(selector).all()
            for i,(row,anchor) in enumerate(zip(native,anchors),first):
                row={**row,'target':urljoin(page.url,anchor.get_attribute('href'))}
                for key in (row['target'],row.get('DOI'),row.get('ProductCode') if row['ContentType']=='B' else None):
                    if key:
                        self.known[key.casefold()]=(row,proof)
                if i in selected_indices:
                    selected.append(row)
            search_url=page.url
            if export=='html':
                html_pages.append(self.store.put(clean_html(page.content(),self.store.sanitizer),'text/html',self.source,page.url,role='search_page',first=first,total=total))
            if issue:
                errors.append(issue);resume_start=issue['resume_start'];incomplete=True
                break
            if first+len(native)>=min(start+limit,total):
                break
            page_number=max(page_number+1,start//payload['RowsPerPage']+1)
            parsed=urlsplit(page.url)
            params=parse_qs(parsed.query)
            params['PageNumber']=[str(page_number)]
            try:
                with page.expect_response(matches,timeout=60000) as pending:
                    if expression:
                        page.evaluate("number => clickPageBtn(number, '')",page_number)
                    else:
                        goto_source(page,urlunsplit((parsed.scheme,parsed.netloc,parsed.path,urlencode(params,doseq=True),'')),self.profile.hosts,wait_until='domcontentloaded',timeout=45000)
                response=pending.value
            except Exception as error:
                resume_start=(page_number-1)*payload['RowsPerPage']
                errors.append({'status':'error','code':'search_response_not_received',
                               'reason':'AMS did not provide the next native search page',
                               'page_number':page_number,'resume_start':resume_start,
                               'error_type':type(error).__name__})
                incomplete=True
                break
        artifact=self.store.put(encoded(selected),'application/json',self.source,search_url,role='records',query=query,options=options,start=start,total=total)
        exported=artifact
        if export=='native':
            exported={'artifacts':evidence,'scope':page_export_scope('native',incomplete)}
        elif export=='html':
            exported={'artifacts':html_pages,'scope':page_export_scope('html',incomplete)}
        elif export=='csv':
            buffer=io.StringIO(newline='')
            fields=list(dict.fromkeys(k for row in selected for k in row))
            writer=csv.DictWriter(buffer,fieldnames=fields)
            writer.writeheader()
            writer.writerows({k:json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else v for k,v in row.items()} for row in selected)
            exported=self.store.put(buffer.getvalue(),'text/csv',self.source,search_url,role='mechanical_export',parent_ref=artifact['ref'])
        elif export=='xlsx':
            exported=self.export_products(selected)
        next_start=resume_start if incomplete else start+len(selected)
        more=next_start<total
        status='partial' if incomplete else 'ok' if selected else 'not_found'
        if export=='xlsx' and selected and exported['status']!='available':status='partial'
        return {'status':status,'source':self.source,'query':query,'options':options,'start':start,'requested':limit,
                'received':len(selected),'total':total,'has_more':more,
                'records':artifact,'export':exported,'evidence':evidence,'preview':[{k:row.get(k) for k in ('Title','Subtitle','ProductCode','DOI','target')} for row in selected[:3]],
                **({'errors':errors,'incomplete_results':True} if incomplete else {}),
                **({'next':{'start':next_start}} if more else {'next':None})}

    def fetch(self,target,*,parts):
        if set(parts)-{'metadata','abstract','fulltext','supplementary','export'}:
            raise ServiceError('invalid_parts','Unknown requested content')
        key=target
        explicit_code=False
        if target.startswith(('http://','https://')):
            parsed=urlsplit(target)
            if parsed.hostname in self.profile.hosts and parsed.path.lower()=='/view':
                values=parse_qs(parsed.query).get('ProductCode',[])
                if len(values)==1:
                    key=values[0]
                    explicit_code=True
        code=explicit_code or re.fullmatch(r'[A-Za-z][A-Za-z0-9-]*/[\w./-]+',key) or key.isupper() and re.fullmatch(r'[A-Z][A-Z0-9_-]*',key)
        if code and key.casefold() not in self.known:
            found=self.search(key,options={'field':'ProductCode'},limit=2)
            if found['total']!=1:
                return {**found,'status':'selection_required' if found['received'] else 'not_found'}
        if key.casefold() not in self.known:
            return super().fetch(target,parts=parts)
        row,proof=self.known[key.casefold()]
        artifact=self.store.put(encoded([row]),'application/json',self.source,row['target'],role='records',evidence_ref=proof['ref'])
        components={}
        for part in parts:
            if part in ('metadata','export'):
                components[part]={'status':'available','artifact':artifact,'format':'native_product_json'}
            elif part=='abstract':
                field='Abstract' if row.get('Abstract') else 'ProductDescription' if row.get('ProductDescription') else None
                components[part]={'status':'available','artifact':artifact,'field':field,'kind':'publisher_abstract' if field=='Abstract' else 'publisher_product_description'} if field else {'status':'not_found','reason':'This native AMS product record contains neither Abstract nor ProductDescription'}
        document_parts=[part for part in parts if part in ('fulltext','supplementary')]
        if document_parts:
            try:
                obtained=super().fetch(row['target'],parts=document_parts)
                components.update(obtained.get('components',{}))
            except ServiceError as error:
                components.update({part:error.result for part in document_parts})
            except Exception as error:
                components.update({part:{'status':'error','code':'document_retrieval_failed','error_type':type(error).__name__} for part in document_parts})
        return {'status':'ok' if all(part['status']=='available' for part in components.values()) else 'partial',
                'source':self.source,'target':target,'records':artifact,'components':components,'evidence':[proof]}
